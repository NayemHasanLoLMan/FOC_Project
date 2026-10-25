"""Development-only Jordan readiness checks; never reads or writes live jobs/invoices."""
import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
import sys
import tempfile

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from development.analysis.compare_invoices import normalize, read_invoice
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.pipeline_config import load_config
from src.foc_extractor.workbook_verification import verify_workbook


def template_checks(path):
    book = load_workbook(path)
    try:
        sheet = book.active
        labels = defaultdict(list)
        for row in sheet:
            for cell in row:
                label = normalize(cell.value)
                if label.startswith(('TOTAL VALUE:', 'COUNTRY OF ULTIMATE DESTINATION',
                                     'DESCRIPTION OF GOODS', 'INVOICE ONLY FOR CUSTOMS PURPOSE')):
                    labels[label].append(cell.coordinate)
        return {'print_area': str(sheet.print_area),
                'duplicate_labels': {label: cells for label, cells in labels.items() if len(cells) > 1}}
    finally:
        book.close()


def writer_probe(source, template, rows=25):
    """Exercise layout capacity independently of OCR accuracy and sample prices."""
    with pymupdf.open(source) as document:
        pages = len(document)
    data = FOCData(source_file=source.name, template_type='JORDAN', pages_total=pages,
                   pages_processed=pages, total_value=rows,
                   items=[LineItem(f'QA ITEM {i}', 1, 'PCS', amount=1) for i in range(rows)])
    with tempfile.TemporaryDirectory() as directory:
        output = Path(directory)/'probe.xlsx'
        populate_template(template, output, data)
        book = load_workbook(output)
        try:
            printed = sum(str(cell.value or '').startswith('QA ITEM ')
                          for row in book.active for cell in row)
        finally:
            book.close()
        error = None
        try:
            verify_workbook(output, data.to_dict(), source, 'JORDAN')
        except ValueError as exc:
            error = str(exc)
    return {'synthetic_requested_rows': rows, 'written_rows': printed,
            'publication_check_passed': error is None, 'publication_error': error}


def readiness(root=ROOT):
    contracts = json.loads((root/'development/analysis/baseline.json').read_text(encoding='utf-8'))['JORDAN']['samples']
    samples = []
    for name, contract in contracts.items():
        source = root/'examples/jordan'/name
        reference = root/'expected_outputs/jordan'/name.replace('-D.pdf', '-I.xlsx')
        with pymupdf.open(source) as document:
            pages = len(document)
            text_pages = sum(bool(page.get_text().strip()) for page in document)
        samples.append({'source': name, 'pages': pages, 'embedded_text_pages': text_pages,
                        'contract_matches_source': pages == contract['pages'] == len(contract['page_checks'])
                            and hashlib.sha256(source.read_bytes()).hexdigest() == contract['source_sha256'],
                        'reference_unchanged': hashlib.sha256(reference.read_bytes()).hexdigest() == contract['reference_sha256'],
                        'manual_goods_rows_for_context_only': len(read_invoice(reference)['items']),
                        'destination_conflict_pages': contract['destination_conflict_pages']})
    template = root/'templates/JORDAN_template.xlsx'
    return {'purpose': 'Jordan development readiness, not an extraction accuracy score',
            'samples': samples, 'source_pages': sum(sample['pages'] for sample in samples),
            'template': template_checks(template),
            'writer_probe': writer_probe(root/'examples/jordan'/next(iter(contracts)), template),
            'automatic_intake_enabled': load_config(root/'pipeline_config.json').jordan_enabled,
            'page_policy': 'Include all pages; use the shipment cover destination; retain conflicting recipients as evidence',
            'remaining_decisions': ['Jordan category/unit rates and total rounding are not approved'],
            'not_tested': ['fresh OCR/vision accuracy', 'native Excel print', 'unseen Jordan PDFs']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', type=Path, default=ROOT/'reports/jordan-readiness/readiness.json')
    args = parser.parse_args()
    result = readiness()
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(result, indent=2))
