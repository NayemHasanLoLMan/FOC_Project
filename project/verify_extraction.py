"""Read-only comparison. Expected workbooks are never imported by extraction."""
import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent
UNITS = {'YDS', 'YD', 'PCS', 'PC', 'SETS', 'SET', 'GROSS', 'GRS', 'G', 'CONES', 'MTRS', 'ROLLS', 'PACKS'}


def normalize(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().upper()


def unit(value):
    return {'PC':'PCS', 'SET':'SETS', 'YD':'YDS', 'GRS':'GROSS'}.get(normalize(value), normalize(value))


def description(value):
    text = re.sub(r'\bFABRICS? SWATCH(?:ES)?\b', 'FABRIC SWATCHES', normalize(value))
    text = re.sub(r'\s*\(\s*', ' (', text)
    text = re.sub(r'\s*\)', ')', text)
    return re.sub(r'\b(WOVEN|KNIT) FABRIC\b', r'\1 FABRICS', text)


def numeric(value):
    if isinstance(value, (int, float)):
        return float(value)
    match = re.fullmatch(r'(?:USD\s*)?(\d+(?:\.\d+)?)', normalize(value))
    return float(match[1]) if match else None


def read_invoice(path):
    """Inspect the actual invoice sheet, independently of its audit."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    rows = {r: [(c, ws.cell(r,c).value) for c in range(1,ws.max_column+1)
                if ws.cell(r,c).value is not None] for r in range(1,ws.max_row+1)}
    result = {'items': [], 'path':str(path)}
    header = None
    for r, cells in rows.items():
        for c, value in cells:
            label = normalize(value)
            right = [v for col,v in cells if col>c and normalize(v)]
            if label.startswith('CUSTOMS INVOICE NO'):
                result['invoice_number'] = next((re.sub(r'\s','',normalize(v)) for v in right if re.match(r'S\s*/',str(v))),None)
            elif label in ('TOTAL:', 'TOTAL :'):
                neighbor=ws.cell(r,c+1).value
                result['package_count'] = neighbor if isinstance(neighbor,(int,float)) else None
            elif label in ('GW:', 'G.W.', 'GW', 'G.W'):
                result['gross_weight_kg'] = next((v for v in right if isinstance(v,(int,float))),None)
            elif label == 'TOTAL VALUE:':
                result['total_value'] = next((numeric(v) for v in right if numeric(v) is not None),None)
            elif label.startswith('CONSIGNEE'):
                result['consignee'] = next((normalize(v) for v in right if 'FOC' not in normalize(v) and 'C/NO' not in normalize(v)),None)
            elif 'DATE OF EXPORTATION' in label:
                for v in right:
                    if isinstance(v,(datetime,date)):
                        result['export_date']=v.strftime('%Y-%m-%d');break
                    if isinstance(v,str) and re.match(r'\d{4}-\d{2}-\d{2}',v):
                        result['export_date']=v[:10];break
                    if isinstance(v,str) and v.startswith('=TODAY'):
                        result['date_is_volatile']=True
            if re.fullmatch(r'[A-Z]+-FOC',label):result['shipping_marks']=label
            if 'DESCRIPTION OF GOODS' in label:header=r
    if header:
        for r,cells in rows.items():
            if r<=header:continue
            if any('TOTAL VALUE' in normalize(v) or 'NO COMMERCIAL VALUE' in normalize(v) for c,v in cells):break
            desc=ws.cell(r,2).value
            if not normalize(desc) or any(t in normalize(desc) for t in ['TOTAL VALUE','COMMERCIAL VALUE','CUSTOMS PURPOSE']):continue
            for c,v in cells:
                u=ws.cell(r,c+1).value
                # Keep invalid numeric units in the baseline denominator.
                is_unit=normalize(u) in UNITS or (c==7 and isinstance(u,str) and u.isnumeric())
                if c>=5 and isinstance(v,(int,float)) and is_unit:
                    extra=ws.cell(r,3).value
                    result['items'].append({'description':description(str(desc)+(' '+str(extra) if extra else '')),
                                            'quantity':float(v),'unit':unit(u),'row':r})
                    break
            else:
                extra=ws.cell(r,3).value
                result['items'].append({'description':description(str(desc)+(' '+str(extra) if extra else '')),
                                        'quantity':numeric(ws.cell(r,7).value),'unit':unit(ws.cell(r,8).value),'row':r})
        if 'total_value' not in result and result['items']:
            last=result['items'][-1]['row']
            cols=[c for r,cells in rows.items() for c,v in cells if normalize(v)=='VALUE' and r<=header+2]
            totals=[ws.cell(r,c).value for r in range(last+1,ws.max_row+1) for c in cols
                    if isinstance(ws.cell(r,c).value,(int,float))]
            if len(totals)==1:result['total_value']=float(totals[0])
    if 'Extraction_Audit' in wb:
        audit={str(r[0]):r[1] for r in wb['Extraction_Audit'].iter_rows(values_only=True) if r[0]}
        result['warnings']=audit.get('warnings')
    if 'Extracted_Items' in wb:
        audit_rows=[r for r in wb['Extracted_Items'].iter_rows(min_row=2,values_only=True) if r[1]]
        result['audit_item_count']=len(audit_rows)
        result['audit_matches_invoice']=[(description(r[1]),numeric(r[2]),unit(r[3])) for r in audit_rows]==[
            (i['description'],i['quantity'],i['unit']) for i in result['items']]
    else:result['audit_matches_invoice']=not result['items']
    wb.close()
    return result


def item_key(item):
    return item['description'], round(item['quantity'],4) if item['quantity'] is not None else None, item['unit']


def quantity_totals(items):
    totals=Counter()
    for item in items:
        if item['quantity'] is not None:totals[item['unit']]+=item['quantity']
    return {u:round(q,4) for u,q in sorted(totals.items())}


def compare(output_dir, expected_dir):
    expected={p.name:read_invoice(p) for p in sorted(expected_dir.glob('*.xlsx'))}
    manifest_path=output_dir/'run_manifest.json'
    manifest=json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else None
    written={Path(r['output']).name for r in manifest if r['status']=='written'} if manifest is not None else None
    generated={p.name:read_invoice(p) for p in sorted(output_dir.glob('*.xlsx')) if written is None or p.name in written}
    fields=['invoice_number','shipping_marks','consignee','package_count','gross_weight_kg','total_value','export_date']
    stats={f:{'matched':0,'total':0} for f in fields}
    cases=[]; used=set(); matches=expected_count=generated_count=0
    quantity_matches=quantity_expected=quantity_generated=0
    for name,g in generated.items():
        target=name if name in expected else None
        if target is None and g.get('invoice_number'):
            candidates=[n for n,e in expected.items() if e.get('invoice_number')==g['invoice_number']]
            if len(candidates)==1:target=candidates[0]
        if target is None or target in used:
            cases.append({'file':name,'matched_reference':None,'generated':g});continue
        used.add(target);e=expected[target];scores={}
        for f in fields:
            if e.get(f) is None:continue
            a,b=e.get(f),g.get(f)
            ok=abs(a-b)<1e-6 if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
            stats[f]['total']+=1;stats[f]['matched']+=int(ok);scores[f]=bool(ok)
        ec=Counter(item_key(i) for i in e['items']);gc=Counter(item_key(i) for i in g['items'])
        count=sum((ec & gc).values());matches+=count;expected_count+=sum(ec.values());generated_count+=sum(gc.values())
        eq=Counter((i['quantity'],i['unit']) for i in e['items'] if i['quantity'] is not None and i['unit'] in UNITS)
        gq=Counter((i['quantity'],i['unit']) for i in g['items'] if i['quantity'] is not None and i['unit'] in UNITS)
        quantity_matches+=sum((eq & gq).values());quantity_expected+=len(e['items']);quantity_generated+=len(g['items'])
        cases.append({'file':name,'matched_reference':target,'fields':scores,'item_matches':count,
                      'expected_quantity_totals':quantity_totals(e['items']),
                      'generated_quantity_totals':quantity_totals(g['items']),
                      'expected':e,'generated':g,'missing_items':[list(k)+[v] for k,v in (ec-gc).items()],
                      'extra_or_incorrect_items':[list(k)+[v] for k,v in (gc-ec).items()]})
    return {'field_accuracy':stats,'item_accuracy':{'matched':matches,'expected':expected_count,'generated':generated_count,
            'recall':matches/expected_count if expected_count else 0,'precision':matches/generated_count if generated_count else 0},
            'quantity_unit_overlap':{'matched':quantity_matches,'expected':quantity_expected,'generated':quantity_generated},
            'generated_files':len(generated),'reference_files':len(expected),
            'incomplete_item_rows':sum(i['quantity'] is None or i['unit'] not in UNITS for g in generated.values() for i in g['items']),
            'matched_files':len(used),'missing_expected_files':sorted(set(expected)-used),
            'incomplete_sources':[r for r in manifest or [] if r['status']!='written'],'cases':cases}


def write_report(result, destination, baseline=None):
    destination.mkdir(parents=True,exist_ok=True)
    (destination/'accuracy.json').write_text(json.dumps({'current':result,'baseline':baseline},indent=2,default=str),encoding='utf-8')
    lines=['# BD sample accuracy\n\n','Manual references are evaluation inputs only. Scores use the actual invoice sheet. ',
           'Items must match description, quantity and unit after limited whitespace/plural normalization. ',
           'No fuzzy overlap is counted as a match. TODAY() reference dates are unscored. ',
           'This is development-set accuracy, not a held-out production estimate.\n\n',
           '| Field | Previous | Current |\n|---|---:|---:|\n']
    for name,st in result['field_accuracy'].items():
        old=baseline['field_accuracy'][name] if baseline else None
        previous=f"{old['matched']}/{old['total']}" if old else '—'
        lines.append(f"| {name} | {previous} | {st['matched']}/{st['total']} |\n")
    for label,data in [('Previous',baseline),('Current',result)]:
        if data:
            a=data['item_accuracy']
            lines.append(f"\n{label} exact item matches: {a['matched']}/{a['expected']} expected ({a['recall']:.1%} recall); {a['matched']}/{a['generated']} generated ({a['precision']:.1%} precision).\n")
    lines.append(f"\nMatched files: {result['matched_files']}. Missing expected outputs: {', '.join(result['missing_expected_files'])}.\n")
    q=result['quantity_unit_overlap']
    lines.append(f"\nQuantity/unit multiset overlap: {q['matched']}/{q['expected']}. This ignores descriptions and does not prove correct row allocation. Incomplete generated item rows: {result['incomplete_item_rows']} (included in item precision denominator for paired files).\n")
    if result['incomplete_sources']:
        lines.append('\nIncomplete generation: '+', '.join(r['source']+' ('+r['status']+')' for r in result['incomplete_sources'])+'. Stale workbooks are excluded.\n')
    lines.append('\n| Output | Reference | Exact items | Invoice/audit parity | Mismatched fields |\n|---|---|---:|---|---|\n')
    for c in result['cases']:
        if not c['matched_reference']:
            lines.append(f"| {c['file']} | Unmatched | Unscored | {c['generated'].get('audit_matches_invoice')} | Unscored |\n");continue
        lines.append(f"| {c['file']} | {c['matched_reference']} | {c['item_matches']}/{len(c['expected']['items'])} | {c['generated'].get('audit_matches_invoice')} | {', '.join(f for f,v in c['fields'].items() if not v)} |\n")
    for c in result['cases']:
        if not c['matched_reference']:continue
        lines.append(f"\n## {c['file']}\n\n")
        lines.append(f"Quantity totals by unit: expected `{c['expected_quantity_totals']}`; generated `{c['generated_quantity_totals']}`. Matching totals alone do not prove correct goods or allocation.\n\n")
        for f,ok in c['fields'].items():
            if not ok:lines.append(f"- {f}: expected `{c['expected'].get(f)}`, generated `{c['generated'].get(f)}`.\n")
        for label,key in [('Missing or incorrect expected items','missing_items'),('Extra or incorrect generated items','extra_or_incorrect_items')]:
            lines.append(f'\n{label}:\n\n')
            for desc,qty,u,count in c[key]:lines.append(f'- {desc}: '+(f'{qty:g}' if qty is not None else '(blank)')+f' {u}'+(f' ({count} rows)' if count>1 else '')+'\n')
            if not c[key]:lines.append('None.\n')
        lines.append(f"\nExtraction warnings: {c['generated'].get('warnings')}.\n")
    (destination/'accuracy.md').write_text(''.join(lines),encoding='utf-8')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=ROOT/'output')
    parser.add_argument('--expected-dir',type=Path,default=ROOT/'expected_outputs')
    parser.add_argument('--report-dir',type=Path,default=ROOT/'reports')
    parser.add_argument('--baseline-dir',type=Path,default=ROOT/'reports/baseline')
    args=parser.parse_args()
    result=compare(args.output_dir,args.expected_dir)
    baseline=compare(args.baseline_dir,args.expected_dir) if args.baseline_dir.exists() else None
    write_report(result,args.report_dir,baseline)
    print(json.dumps({k:v for k,v in result.items() if k!='cases'},indent=2))
