from datetime import datetime
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.workbook_verification import verify_workbook
from src.foc_extractor.invoice_preview import invoice_sheet


class JordanWriterTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.source = self.root/'unseen.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page(); pdf.save(self.source)

    def data(self, count=1):
        return FOCData(source_file=self.source.name, template_type='JORDAN', pages_total=1,
                       pages_processed=1, export_date='2027-02-23', invoice_number='S/NEW&01/27',
                       consignee='NEW RECIPIENT', consignee_address='ADDRESS ONE\nADDRESS TWO',
                       shipping_marks='NEW-FOC', destination_city='AMMAN', destination_country='Jordan',
                       carton_range='J01-J03, J05', package_count=4, gross_weight_kg=5.125,
                       total_value=0 if count == 0 else count * 0.25,
                       items=[LineItem('GOODS '+str(i), 1.125, 'YDS', weight_kg=0.02, amount=0.25,
                                       unit_price=0.2222) for i in range(count)])

    def write(self, data):
        path = self.root/'invoice.xlsx'
        populate_template(ROOT/'templates/JORDAN_template.xlsx', path, data)
        verify_workbook(path, data.to_dict(), self.source, 'JORDAN')
        return path

    def test_blank_master_has_one_section_per_label_and_valid_print_metadata(self):
        master = ROOT/'templates/JORDAN_template.xlsx'
        book = load_workbook(master); sheet = book.active
        for label in ('TOTAL VALUE:', 'DESCRIPTION OF GOODS', 'COUNTRY OF ULTIMATE DESTINATION:',
                      'NO COMMERCIAL VALUE', 'INVOICE ONLY FOR CUSTOMS PURPOSE'):
            self.assertEqual(sum(cell.value == label for row in sheet for cell in row),1,label)
        self.assertIn('$A$1:$J$39',sheet.print_area)
        self.assertEqual((sheet.page_setup.fitToWidth,sheet.page_setup.fitToHeight),(1,0))
        self.assertFalse(any(cell.data_type=='f' for row in sheet for cell in row))
        book.close()
        with zipfile.ZipFile(master) as archive:
            xml=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(len(xml.findall('{*}pageMargins')),1)

    def test_zero_one_fifteen_seventeen_twentyfive_and_sixty_rows_without_loss(self):
        for count in (0, 1, 15, 17, 25, 60):
            with self.subTest(count=count):
                data = self.data(count); path = self.write(data)
                book = load_workbook(path); sheet = book.active
                totals = [c for row in sheet for c in row if c.value == 'TOTAL VALUE:']
                self.assertEqual(len(totals), 1)
                self.assertGreater(totals[0].row, 26 + count)
                self.assertEqual(sheet.cell(totals[0].row, 9).value, data.total_value)
                self.assertEqual(sheet.print_title_rows, '$24:$24')
                self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 0))
                self.assertEqual(sheet.oddFooter.right.text, 'Page &P of &N')
                with zipfile.ZipFile(path) as archive:
                    xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
                    header = xml.find('{*}headerFooter/{*}oddHeader').text
                    self.assertIn('S/NEW&&01/27', header)
                self.assertEqual([sheet.cell(27+i, 2).value for i in range(count)], [it.description for it in data.items])
                self.assertTrue(all(sheet.cell(27+i, 2).border.left.style == 'thin' for i in range(max(1,count))))
                self.assertEqual(sheet['C9'].value, datetime(2027, 2, 23))
                self.assertEqual(sheet['C9'].alignment.horizontal, 'left')
                occupied = set()
                for merged in sheet.merged_cells.ranges:
                    self.assertFalse(occupied.intersection(merged.cells))
                    occupied.update(merged.cells)
                book.close()

    def test_blank_fields_and_zero_values_remain_distinct_and_no_origin_is_invented(self):
        data = FOCData(source_file=self.source.name, template_type='JORDAN', pages_total=1, pages_processed=1,
                       gross_weight_kg=0, total_value=0, items=[LineItem('SAMPLE',0,None,weight_kg=0,amount=0)])
        path = self.write(data); book = load_workbook(path); sheet = book.active
        for address in ('C9','C11','I9','I10','I11','D16','C18','B21','C12','C15'):
            self.assertIsNone(sheet[address].value,address)
        for address in ('F21','F27','I27'):
            self.assertEqual(sheet[address].value,0,address)
        self.assertEqual(sheet['G27'].value,'0')
        self.assertIsNone(sheet['H27'].value)
        self.assertFalse(any('CHINA' in str(c.value) or 'TODAY(' in str(c.value) for row in sheet for c in row))
        book.close()

    def test_long_text_wraps_and_all_address_lines_survive(self):
        data = self.data()
        data.consignee = 'LONG RECIPIENT NAME ' * 6
        data.consignee_address = '\n'.join('ADDRESS LINE '+str(i) for i in range(8))
        data.items[0].description = 'COTTON FABRIC WITH A READABLE LONG DESCRIPTION ' * 10
        path = self.write(data); book = load_workbook(path); sheet = book.active
        self.assertEqual(sheet['C15'].value, '\n'.join(data.consignee_address.splitlines()[3:]))
        self.assertGreater(sheet.row_dimensions[11].height,26)
        self.assertGreater(sheet.row_dimensions[27].height,26)
        self.assertTrue(sheet['B27'].alignment.wrap_text)
        book.close()

    def test_quantity_and_unit_share_one_printed_column_and_keep_raw_data(self):
        data=self.data(6)
        cases=[(90,'YDS','90 YDS'),(37,'PCS','37 PCS'),(1.125,'SETS','1.125 SETS'),
               (0,'PCS','0 PCS'),(None,'YDS','YDS'),(None,None,None)]
        for item,(quantity,unit,_) in zip(data.items,cases):item.quantity=quantity;item.unit=unit
        path=self.write(data);book=load_workbook(path);sheet=book.active
        self.assertIn('G24:H24',sheet.merged_cells)
        self.assertEqual(sheet['G24'].value,'QUANTITY');self.assertIsNone(sheet['H24'].value)
        for index,(quantity,unit,label) in enumerate(cases):
            row=27+index
            self.assertIn(f'G{row}:H{row}',sheet.merged_cells)
            self.assertEqual(sheet.cell(row,7).value,label)
            self.assertIsNone(sheet.cell(row,8).value)
            self.assertEqual(book['Extracted_Items'].cell(index+2,3).value,quantity)
            self.assertEqual(book['Extracted_Items'].cell(index+2,4).value,unit)
        book.close()
        cells=[cell for row in invoice_sheet(path.read_bytes())['rows'] for cell in row['cells']]
        self.assertTrue(any(cell['colspan']==2 and cell['parts'][0]['text']=='90 YDS' for cell in cells))
        self.assertFalse(any(part['text']=='UNIT' for cell in cells for part in cell['parts']))
        for action in ('wrong unit','unmerged'):
            path=self.write(data);book=load_workbook(path);sheet=book.active
            if action=='wrong unit':sheet['G27']='90 PCS'
            else:sheet.unmerge_cells('G27:H27')
            book.save(path);book.close()
            with self.assertRaises(ValueError):verify_workbook(path,data.to_dict(),self.source,'JORDAN')

    def test_retained_separate_quantity_and_unit_layout_still_verifies(self):
        data=self.data();path=self.write(data);book=load_workbook(path);sheet=book.active
        sheet.unmerge_cells('G24:H24');sheet['H24']='UNIT'
        for row in range(27,33):sheet.unmerge_cells(f'G{row}:H{row}')
        sheet['G27']=data.items[0].quantity;sheet['G27'].number_format='#,##0.####'
        sheet['H27']=data.items[0].unit
        book.save(path);book.close()
        verify_workbook(path,data.to_dict(),self.source,'JORDAN')

    def test_excessive_text_refuses_before_overwriting_existing_output(self):
        path = self.root/'invoice.xlsx'; path.write_bytes(b'previous invoice')
        data = self.data(); data.items[0].description = 'UNPRINTABLE ' * 2000
        with self.assertRaisesRegex(ValueError,'printable height'):
            populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
        self.assertEqual(path.read_bytes(), b'previous invoice')

    def test_formula_like_source_text_is_literal_and_source_currency_is_preserved(self):
        data = self.data(); data.items[0].description = '=1+2'; data.invoice_number = '=9+9'; data.currency='EUR'
        path = self.write(data); book = load_workbook(path); sheet = book.active
        for address in ('B27','C18'):
            self.assertEqual(sheet[address].data_type,'s')
        self.assertEqual(sheet['I24'].value,'VALUE (EUR)')
        self.assertEqual(sheet['I35'].number_format,'"EUR "#,##0.00')
        book.close()

    def test_independent_verifier_rejects_altered_header_goods_total_format_visibility_and_declarations(self):
        for coordinate, value in (('C11','WRONG'), ('I11','WRONG'), ('F21',999), ('I35',999),
                                  ('G27',999), ('F27',999), ('A38',None)):
            with self.subTest(coordinate=coordinate):
                data=self.data();path=self.write(data);book=load_workbook(path)
                book.active[coordinate]=value;book.save(path);book.close()
                with self.assertRaises(ValueError):verify_workbook(path,data.to_dict(),self.source,'JORDAN')
        for change in ('format','hidden','print','date'):
            data=self.data();path=self.write(data);book=load_workbook(path);sheet=book.active
            if change=='format':sheet['I35'].number_format=';;;'
            elif change=='hidden':sheet.row_dimensions[27].hidden=True
            elif change=='print':sheet.print_area='A24:J39'
            else:sheet['C9'].number_format='0'
            book.save(path);book.close()
            with self.assertRaises(ValueError):verify_workbook(path,data.to_dict(),self.source,'JORDAN')
