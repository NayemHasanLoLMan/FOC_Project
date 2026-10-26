from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.workbook_verification import verify_workbook


class ShipmentWordingTests(unittest.TestCase):
    def test_destination_uses_available_fields_without_route_or_transport_defaults(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'unseen.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page(); pdf.save(source)
            for city, country, expected in ((None, None, ''), (None, 'Bangladesh', 'Bangladesh'),
                                             ('NEW CITY, BD', 'Bangladesh', 'NEW CITY, BD')):
                for count in (1, 9):
                    with self.subTest(city=city, country=country, count=count):
                        data = FOCData(source_file=source.name, template_type='BD', pages_processed=1, pages_total=1,
                                       destination_city=city, destination_country=country,
                                       items=[LineItem('FABRIC', 1, 'YDS') for _ in range(count)])
                        output = Path(directory)/'invoice.xlsx'
                        populate_template(ROOT/'templates/BD_template.xlsx', output, data)
                        verify_workbook(output, data.to_dict(), source)
                        book = openpyxl.load_workbook(output)
                        self.assertEqual(book.active['A18'].value, 'SHIPMENT TO: '+expected)
                        book.close()

    def test_old_custom_template_cannot_restore_fixed_route(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory)/'BD_template.xlsx'
            book = openpyxl.load_workbook(ROOT/'templates/BD_template.xlsx')
            self.assertEqual(book.active['A18'].value, 'SHIPMENT TO:')
            book.active['A18'] = 'SHIPMENT FROM HONG KONG TO DHAKA BY AIR'
            book.save(template); book.close()
            output = Path(directory)/'invoice.xlsx'
            populate_template(template, output, FOCData(source_file='new.pdf', template_type='BD'))
            book = openpyxl.load_workbook(output)
            self.assertEqual(book.active['A18'].value, 'SHIPMENT TO: ')
            book.close()

    def test_publication_rejects_stale_or_changed_route(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'new.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page(); pdf.save(source)
            data = FOCData(source_file=source.name, template_type='BD', pages_processed=1, pages_total=1,
                           destination_country='Bangladesh', items=[LineItem('FABRIC', 1, 'YDS')])
            output = Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx', output, data)
            for value in ('SHIPMENT FROM HONG KONG TO DHAKA BY AIR', 'SHIPMENT TO: OTHER CITY', '="SHIPMENT TO: Bangladesh"'):
                book = openpyxl.load_workbook(output)
                book.active['A18'] = value
                book.save(output); book.close()
                with self.assertRaisesRegex(ValueError, 'Bangladesh shipment destination'):
                    verify_workbook(output, data.to_dict(), source)


if __name__ == '__main__':
    unittest.main()
