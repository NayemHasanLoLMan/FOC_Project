"""Presentation changes preserve source data and historical publication checks."""
from pathlib import Path
import sys
import tempfile
import unittest

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.workbook_verification import verify_workbook


class PresentationTests(unittest.TestCase):
    def test_bangladesh_never_adds_price_columns_and_preserves_calculation_data(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary); source = root/'new.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page(); pdf.save(source)
            for mark in ('ROSE-FOC', 'GTL-FOC', None):
                data = FOCData(source_file=source.name, template_type='BD', pages_processed=1, pages_total=1,
                               consignee='THE ROSE DRESSES LTD', shipping_marks=mark,
                               carton_range='A01-A02', total_value=2,
                               items=[LineItem('DIE MOULD', 2, 'SETS', unit_price=1, amount=2)])
                before = data.to_dict(); path = root/'invoice.xlsx'
                populate_template(ROOT/'templates/BD_template.xlsx', path, data)
                verify_workbook(path, data.to_dict(), source)
                book = load_workbook(path); sheet = book.active
                self.assertEqual((sheet['G30'].value, sheet['G33'].value, sheet['H33'].value), ('QUANTITY',2,'SETS'))
                self.assertFalse(any('UNIT PRICE' in str(c.value) or c.value=='VALUE (USD)' for row in sheet for c in row))
                self.assertEqual(book['Extraction_Audit'].sheet_state, 'hidden')
                self.assertEqual(book['Extracted_Items'].sheet_state, 'hidden')
                self.assertEqual((book['Extracted_Items']['G2'].value, book['Extracted_Items']['H2'].value), (1,2))
                self.assertEqual(sheet['G40'].value,2)
                self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
                self.assertFalse(sheet.page_setup.usePrinterDefaults)
                self.assertEqual(data.to_dict(), before)
                book.close()

    def test_jordan_common_footer_mixed_origins_and_retained_old_layout(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary); source = root/'new.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page(); pdf.save(source)
            data = FOCData(source_file=source.name, template_type='JORDAN', pages_processed=1, pages_total=1,
                           items=[LineItem('LABEL',2,'PCS',origin_country='CHINA'),
                                  LineItem('THREAD',1,'CONES',origin_country='CHINA')])
            path = root/'invoice.xlsx'
            populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
            verify_workbook(path,data.to_dict(),source,'JORDAN')
            book=load_workbook(path);sheet=book.active
            self.assertEqual(sheet['A36'].value,'COUNTRY OF ORIGIN: CHINA')
            self.assertEqual(sheet['B27'].value,'LABEL')
            self.assertTrue(sheet.column_dimensions['K'].hidden)
            sheet['A36']=None
            for index,item in enumerate(data.items,27):
                sheet.cell(index,2,item.description+'\nCOUNTRY OF ORIGIN: CHINA')
            book.save(path);book.close()
            verify_workbook(path,data.to_dict(),source,'JORDAN')
            data.items[1].origin_country='VIETNAM'
            populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
            verify_workbook(path,data.to_dict(),source,'JORDAN')
            book=load_workbook(path)
            self.assertIsNone(book.active['A36'].value)
            self.assertIn('VIETNAM',book.active['B28'].value)
            book.close()

