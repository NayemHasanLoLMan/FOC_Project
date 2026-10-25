import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from dataclasses import replace
from unittest.mock import patch

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor import pipeline, queue_worker
from src.foc_extractor.job_store import JobStore, SCHEMA_VERSION
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.ocr_engine import OCRResult
from src.foc_extractor.pipeline_config import load_config, attempt_metadata
from src.foc_extractor.regions import REGIONS, select_region
from src.foc_extractor.runtime_settings import editable_settings, settings_config
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.workbook_verification import verify_workbook


class RegionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        config = self.root/'pipeline_config.json'
        config.write_text((ROOT/'pipeline_config.json').read_text())
        self.config = load_config(config)
        self.store = JobStore(self.config.state_dir/'jobs.sqlite3')
        self.pdf = self.root/'shipment.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page(); pdf.save(self.pdf)
        self.digest = hashlib.sha256(self.pdf.read_bytes()).hexdigest()

    def data(self, region):
        return FOCData(source_file=self.pdf.name, template_type=region, pages_processed=1, pages_total=1,
                       destination_country=REGIONS[region]['country'], total_value=2,
                       items=[LineItem('FABRIC', 10, 'YDS', hs_code='00123456', weight_kg=3, amount=2)])

    def printed(self, region):
        data = self.data(region)
        path = self.root/(region+'.xlsx')
        populate_template(ROOT/f'templates/{region}_template.xlsx', path, data)
        return path, data

    def test_schema_five_migration_preserves_jobs_attempt_metadata_and_paths(self):
        job = self.store.enqueue(self.pdf, self.digest, self.pdf)
        attempt = self.store.claim_next('old', {'template_sha256':'historical'})
        with self.store._connection(write=True) as db:
            db.execute('ALTER TABLE jobs DROP COLUMN region')
            db.execute('PRAGMA user_version=5')
        with self.assertRaisesRegex(ValueError,'migrate'):
            JobStore(self.store.path,read_only=True)
        migrated = JobStore(self.store.path)
        self.assertEqual(migrated.get_job(job)['region'],'BD')
        self.assertEqual(migrated.get_job(job)['snapshot_path'],str(self.pdf))
        self.assertEqual(migrated.get_attempt(attempt['id'])['metadata'],attempt['metadata'])
        with migrated._connection() as db:
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0],SCHEMA_VERSION)
            self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0],'ok')

    def test_region_is_durable_and_existing_source_cannot_be_reassigned(self):
        job = self.store.enqueue(self.pdf,self.digest,self.pdf,region='INDIA')
        self.assertEqual(JobStore(self.store.path).get_job(job)['region'],'INDIA')
        with self.assertRaisesRegex(ValueError,'reassigned'):
            self.store.enqueue(self.pdf,self.digest,self.pdf,region='BD')
        with self.assertRaises(ValueError):
            self.store.claim_next('wrong',attempt_metadata(self.config,'BD'))
        attempt=self.store.claim_next('worker',lambda j:attempt_metadata(self.config,j['region']))
        self.assertEqual(attempt['metadata']['region'],'INDIA')
        self.assertEqual(attempt['metadata']['template_sha256'],hashlib.sha256((ROOT/'templates/INDIA_template.xlsx').read_bytes()).hexdigest())

    def test_unknown_region_and_wrong_template_fail_before_ocr(self):
        for region, template in [('OTHER',None),('INDIA',ROOT/'templates/BD_template.xlsx')]:
            with self.subTest(region=region),patch.object(pipeline,'analyze_pdf') as ocr:
                with self.assertRaises(ValueError):
                    pipeline.run(self.pdf,region=region,template=template,output_dir=self.root/'out')
                ocr.assert_not_called()

    def test_destination_mismatch_does_not_write_invoice(self):
        with patch.object(pipeline,'analyze_pdf',return_value=OCRResult(pages=['DESTINATION COUNTRY INDIA'])), \
             patch.object(pipeline,'parse_result',return_value=self.data('INDIA')),patch.object(pipeline,'populate_template') as writer:
            with self.assertRaisesRegex(ValueError,'conflicts'):
                pipeline.run(self.pdf,region='BD',output_dir=self.root/'out',local_model=None)
            writer.assert_not_called()

    def test_cover_conflict_and_unsupported_country_are_not_ignored(self):
        for country in ('India','Nepal'):
            with self.subTest(country=country),self.assertRaisesRegex(ValueError,'conflicts'):
                select_region(self.data('BD'),'BD','DESTINATION COUNTRY: '+country)
        # A later packing recipient is not an instruction to reroute a Jordan consignment.
        self.assertEqual(select_region(self.data('JORDAN'),'JORDAN','DESTINATION COUNTRY JORDAN'),'JORDAN')

    def test_reference_folders_and_children_are_protected_for_both_fields(self):
        for name in ('expected_output','expected_outputs'):
            for suffix in ('','/india'):
                for key in ('incoming_dir','output_dir'):
                    with self.subTest(name=name,suffix=suffix,key=key),self.assertRaisesRegex(ValueError,'reference'):
                        settings_config(self.config,{**editable_settings(self.config),key:str(ROOT/(name+suffix))})
                with self.subTest(name=name,suffix=suffix),patch.object(pipeline,'analyze_pdf') as ocr:
                    with self.assertRaisesRegex(ValueError,'reference'):
                        pipeline.run(self.pdf,output_dir=ROOT/(name+suffix),local_model=None)
                    ocr.assert_not_called()

    def test_verifier_uses_each_regions_actual_columns_and_rejects_tampering(self):
        for region in REGIONS:
            with self.subTest(region=region):
                path,data=self.printed(region)
                verify_workbook(path,data.to_dict(),self.pdf,region)
                book=load_workbook(path);book.active.cell(REGIONS[region]['start'],REGIONS[region]['quantity'],999)
                book.save(path);book.close()
                with self.assertRaisesRegex(ValueError,'Printed goods'):
                    verify_workbook(path,data.to_dict(),self.pdf,region)

    def test_regional_identifier_weight_print_area_and_coverage_are_verified(self):
        for region, col in [('INDIA',5),('JORDAN',6)]:
            path,data=self.printed(region)
            book=load_workbook(path);book.active.cell(REGIONS[region]['start'],col,'wrong')
            book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'Printed goods'):
                verify_workbook(path,data.to_dict(),self.pdf,region)
            path,data=self.printed(region)
            book=load_workbook(path);book.active.print_area='A1:J20';book.save(path);book.close()
            with self.assertRaisesRegex(ValueError,'print area'):
                verify_workbook(path,data.to_dict(),self.pdf,region)
            data.pages_total=2
            with self.assertRaisesRegex(ValueError,'page coverage'):
                verify_workbook(path,data.to_dict(),self.pdf,region)

    def test_jordan_overflow_is_preserved_and_can_be_verified(self):
        data = self.data('JORDAN')
        data.items = [LineItem('ITEM '+str(i), 1, 'PCS') for i in range(17)]
        path = self.root/'JORDAN.xlsx'
        populate_template(ROOT/'templates/JORDAN_template.xlsx', path, data)
        verify_workbook(path, data.to_dict(), self.pdf, 'JORDAN')

    def test_worker_records_actual_regional_template_without_bd_fallback(self):
        job=self.store.enqueue(self.pdf,self.digest,self.pdf,region='INDIA')
        def convert(source, **kwargs):
            self.assertEqual(kwargs['region'],'INDIA')
            self.assertEqual(kwargs['template'],ROOT/'templates/INDIA_template.xlsx')
            path,data=self.printed('INDIA')
            target=kwargs['output_dir']/path.name;target.write_bytes(path.read_bytes())
            (kwargs['output_dir']/(source.stem+'.json')).write_text(json.dumps(data.to_dict()))
            return target
        with patch.object(queue_worker,'run',side_effect=convert),queue_worker.QueueWorker(self.config,self.store) as worker:
            aid=worker.run_once()
        attempt=self.store.get_attempt(aid)
        self.assertEqual(attempt['state'],'completed',attempt['error'])
        self.assertEqual(attempt['metadata']['region'],'INDIA')
        self.assertEqual(self.store.get_job(job)['region'],'INDIA')
        self.assertEqual(json.loads(Path(attempt['receipt_path']).read_text())['region'],'INDIA')

    def regional_convert(self, source, **kwargs):
        region=kwargs['region']
        self.assertEqual(kwargs['template'],ROOT/f'templates/{region}_template.xlsx')
        path,data=self.printed(region)
        target=kwargs['output_dir']/path.name;target.write_bytes(path.read_bytes())
        (kwargs['output_dir']/(source.stem+'.json')).write_text(json.dumps(data.to_dict()))
        return target

    def test_prepared_regional_restart_retains_region_and_original_destination(self):
        self.store.enqueue(self.pdf,self.digest,self.pdf,region='JORDAN')
        with patch.object(queue_worker,'run',side_effect=self.regional_convert),queue_worker.QueueWorker(self.config,self.store) as worker:
            with patch.object(queue_worker,'_publish_copy',side_effect=PermissionError('locked')):
                aid=worker.run_once()
        self.assertEqual(self.store.get_attempt(aid)['state'],'waiting_to_save')
        with patch.object(queue_worker,'run') as convert,queue_worker.QueueWorker(replace(self.config,output_dir=self.root/'other-output'),self.store) as worker:
            worker.run_once()
        convert.assert_not_called()
        attempt=self.store.get_attempt(aid)
        self.assertEqual(attempt['state'],'completed',attempt['error'])
        self.assertEqual(attempt['metadata']['region'],'JORDAN')
        self.assertEqual(Path(attempt['output_path']).parent,self.config.output_dir)

    def test_retry_keeps_job_region_and_fingerprints_actual_template(self):
        job=self.store.enqueue(self.pdf,self.digest,self.pdf,region='INDIA')
        with patch.object(queue_worker,'run',side_effect=ValueError('unreadable')),queue_worker.QueueWorker(self.config,self.store) as worker:
            first=worker.run_once()
        self.store.retry(job)
        with patch.object(queue_worker,'run',side_effect=self.regional_convert),queue_worker.QueueWorker(self.config,self.store) as worker:
            second=worker.run_once()
        self.assertEqual(self.store.get_attempt(first)['state'],'failed')
        attempt=self.store.get_attempt(second)
        self.assertEqual(attempt['state'],'completed',attempt['error'])
        self.assertEqual(attempt['metadata']['region'],'INDIA')
        self.assertEqual(attempt['metadata']['template_sha256'],attempt_metadata(self.config,'INDIA')['template_sha256'])

    def test_direct_regional_failure_leaves_existing_output_untouched(self):
        out=self.root/'out';out.mkdir();target=out/'shipment.xlsx';target.write_bytes(b'existing invoice')
        data=self.data('INDIA');data.items=[LineItem('UNPRINTABLE '+('X '*1800),1,'PCS')]
        with patch.object(pipeline,'analyze_pdf',return_value=OCRResult(pages=['DESTINATION COUNTRY INDIA'])), \
             patch.object(pipeline,'parse_result',return_value=data):
            with self.assertRaises(ValueError):
                pipeline.run(self.pdf,region='INDIA',output_dir=out,local_model=None)
        self.assertEqual(target.read_bytes(),b'existing invoice')
        self.assertEqual(list(out.iterdir()),[target])


if __name__=='__main__':unittest.main()
