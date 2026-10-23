import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import tempfile
import json
import openpyxl
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.excel_writer import populate_template, _next_writable
from src.foc_extractor.pipeline import _resolve_template
from tools.compare_invoices import read_invoice, compare


class WriterTests(unittest.TestCase):
    def test_unseen_invoices_inherit_master_highlights_and_alignment_after_growth(self):
        with tempfile.TemporaryDirectory() as directory:
            for count in (1, 9):
                path=Path(directory)/f'new_shipment_{count}.xlsx'
                data=FOCData(source_file='unseen_shipment.pdf',template_type='BD',
                             invoice_number='S/NEW/27',destination_country='Bangladesh',
                             gross_weight_kg=3.1,net_weight_kg=2.7,total_value=11.5,
                             items=[LineItem(f'NEW MATERIAL {n}',1.25,'YDS') for n in range(count)])
                populate_template(ROOT/'templates/BD_template.xlsx',path,data)
                wb=openpyxl.load_workbook(path);ws=wb.active
                extra=max(0,count-6)
                for coordinate in ['G27','G28',f'G{40+extra}',f'A{43+extra}']:
                    self.assertEqual(ws[coordinate].font.color.rgb[-6:],'FF0000',coordinate)
                for coordinate in ['C25','E21']:
                    self.assertEqual(ws[coordinate].font.color.rgb[-6:],'0000FF',coordinate)
                self.assertEqual(ws['C25'].alignment.horizontal,'left')
                for row in range(33,33+count):
                    self.assertEqual(ws.cell(row,2).alignment.horizontal,'left')
                    self.assertEqual(ws.cell(row,7).alignment.horizontal,'right')
                    self.assertEqual(ws.cell(row,8).alignment.horizontal,'left')
                self.assertEqual(ws[f'G{40+extra}'].alignment.horizontal,'right')
                for row in ws[f'F{40+extra}:I{40+extra}']:
                    for cell in row:
                        self.assertIsNone(cell.fill.patternType)
                        for edge in ('top','bottom','left','right'):
                            self.assertIsNone(getattr(getattr(cell.border,edge),'style',None))
                self.assertEqual(ws['G28'].alignment.horizontal,'right')
                self.assertEqual(read_invoice(path)['total_value'],11.5)
                wb.close()

    def test_wrapped_rows_and_continuation_identity_preserve_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            data=FOCData(source_file='new.pdf',template_type='BD',invoice_number='S/12&Page/26',
                         items=[LineItem('FIRST LINE\nSECOND LINE\nTHIRD LINE',1,'PCS')])
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path);ws=wb.active
            self.assertGreaterEqual(ws.row_dimensions[33].height,45)
            self.assertEqual(ws.row_dimensions[28].height,4)
            self.assertEqual(ws['C25'].value,'S/12&Page/26')
            self.assertEqual(ws.oddHeader.right.text,'S/12&&Page/26')
            self.assertIsNone(ws.firstHeader.right.text)
            self.assertEqual(ws.firstFooter.right.text,'Page &P of &N')
            wb.close()

    def test_footer_text_has_room_after_detail_rows_grow(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'new.xlsx'
            data=FOCData(source_file='new.pdf',template_type='BD',
                         items=[LineItem(f'ITEM {n}',1,'PCS') for n in range(7)])
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path)
            for row in wb.active:
                for cell in row:
                    if cell.value in ('NO COMMERCIAL VALUE','INVOICE ONLY FOR CUSTOMS PURPOSE'):
                        self.assertGreaterEqual(wb.active.row_dimensions[cell.row].height,15)
            wb.close()

    def test_goods_carton_labels_are_separate_and_do_not_replace_goods(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'anonymous.xlsx'
            data=FOCData(source_file='anonymous.pdf',template_type='BD',carton_range='A01-A03, B01',
                         items=[LineItem('CLOTH',12,'YDS')])
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path);ws=wb.active
            self.assertEqual((ws['A33'].value,ws['A34'].value),('A01-A03','B01'))
            self.assertEqual((ws['B33'].value,ws['G33'].value),('CLOTH',12))
            self.assertIsNone(ws['B34'].value)
            wb.close()
            data.carton_range=', '.join(f'A{i:02d}' for i in range(1,8))
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path)
            self.assertEqual(wb.active['A39'].value,'A07')
            self.assertEqual(wb.active['F41'].value,'TOTAL VALUE:')
            wb.close()
            for marks,label in [('ROSE-FOC','4'),('CIPL-FOC','4 PKGS')]:
                data.carton_range=None;data.package_count=4;data.shipping_marks=marks
                populate_template(ROOT/'templates/BD_template.xlsx',path,data)
                wb=openpyxl.load_workbook(path)
                self.assertEqual(wb.active['A33'].value,label)
                wb.close()

    def test_decimal_quantities_and_address_are_printable(self):
        data=FOCData(source_file='new.pdf',template_type='BD',carton_range='A01-A12, B01',
                     gross_weight_kg=1.125,net_weight_kg=1.025,
                     consignee_address='A complete address longer than one narrow spreadsheet cell',
                     items=[LineItem('SHANK',8.125,'G')])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path);ws=wb.active
            self.assertEqual(ws['G33'].value,8.125)
            self.assertEqual(ws['G33'].number_format,'#,##0.####')
            self.assertEqual(ws['G33'].alignment.vertical,'top')
            self.assertEqual(ws['G27'].value,1.125)
            self.assertEqual((ws['F28'].value,ws['G28'].value,ws['H28'].value),(None,None,None))
            self.assertIsNone(read_invoice(path).get('net_weight_kg'))
            self.assertEqual(ws['G27'].number_format,'0.00############')
            self.assertIn('B12:F12',{str(r) for r in ws.merged_cells.ranges})
            self.assertIn('A01-A12, B01',ws['H13'].value)
            self.assertEqual(ws.print_title_rows,'$30:$30')
            wb.close()

    def test_all_items_print_and_total_stays_below(self):
        data=FOCData(source_file='synthetic.pdf',template_type='BD',export_date='2027-01-23',
                     items=[LineItem(f'MATERIAL {i}',i+1,'PCS') for i in range(19)],warnings=['Review'])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            invoice=read_invoice(path)
            self.assertEqual(len(invoice['items']),19)
            self.assertTrue(invoice['audit_matches_invoice'])
            wb=openpyxl.load_workbook(path);ws=wb.active
            self.assertEqual(ws['F53'].value,'TOTAL VALUE:')
            self.assertIsNone(ws['G53'].value)
            self.assertIsNone(ws['A29'].value)
            self.assertIsInstance(ws['C9'].value,__import__('datetime').datetime)
            self.assertIn('56',str(ws.print_area))
            occupied=set()
            for merged in ws.merged_cells.ranges:
                cells=set(merged.cells)
                self.assertFalse(occupied & cells,'Overlapping merged cells can prevent Excel from opening the workbook')
                occupied.update(cells)
            wb.close()

    def test_missing_fields_are_blank_and_source_text_is_literal(self):
        data=FOCData(source_file='=1+1',template_type='BD',consignee='=1+1',
                     consignee_address='\n'.join(f'ADDRESS {i}' for i in range(8)),
                     items=[LineItem('=1+1',unit='PCS')],warnings=['REVIEW REQUIRED'])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path);ws=wb.active
            for coordinate in ['C9','H12','C25','B27','G27','F28','G28','H28','G40','G33','A29']:
                self.assertIsNone(ws[coordinate].value,coordinate)
            for coordinate in ['B11','B33']:
                self.assertEqual(ws[coordinate].value,'=1+1')
                self.assertEqual(ws[coordinate].data_type,'s')
            self.assertIn('ADDRESS 7',ws['B17'].value)
            self.assertEqual(wb['Extracted_Items']['B2'].data_type,'s')
            self.assertFalse(any('REVIEW' in str(c.value) or 'DRAFT' in str(c.value) for row in ws for c in row))
            wb.close()

    def test_merged_label_fallback_avoids_nonanchor(self):
        ws=openpyxl.Workbook().active;ws.merge_cells('A1:B1');ws['A1']='LABEL'
        self.assertEqual(_next_writable(ws,ws['A1']).coordinate,'C1')

    def test_missing_explicit_template_fails(self):
        with self.assertRaises(FileNotFoundError):_resolve_template(Path('missing.xlsx'),'BD')

    def test_invalid_template_fails_before_ocr(self):
        from src.foc_extractor.pipeline import run
        with patch('src.foc_extractor.pipeline.analyze_pdf') as ocr:
            with self.assertRaises(FileNotFoundError):
                run(ROOT/'examples/26S5075-D.pdf',template=Path('missing.xlsx'))
            ocr.assert_not_called()

    def test_verifier_does_not_accept_nonempty_wrong_items(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);expected=root/'expected';output=root/'output'
            expected.mkdir();output.mkdir()
            for folder,qty in [(expected,10),(output,999)]:
                populate_template(ROOT/'templates/BD_template.xlsx',folder/'invoice.xlsx',
                                  FOCData(source_file='a.pdf',items=[LineItem('COTTON',qty,'YDS')],template_type='BD'))
            result=compare(output,expected)
            self.assertEqual(result['item_accuracy']['matched'],0)
            self.assertEqual(result['item_accuracy']['expected'],1)
            (output/'run_manifest.json').write_text(json.dumps([{'source':'a.pdf','status':'failed'}]))
            failed=compare(output,expected)
            self.assertEqual(failed['matched_files'],0)
            self.assertEqual(len(failed['incomplete_sources']),1)
