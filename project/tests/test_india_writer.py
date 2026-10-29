from datetime import datetime
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.workbook_verification import verify_workbook
from src.foc_extractor.ocr_engine import OCRResult
from src.foc_extractor.pipeline import run


class IndiaWriterTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.pdf = self.root/'unseen.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page();pdf.save(self.pdf)

    def data(self,count=1):
        return FOCData(source_file=self.pdf.name,template_type='INDIA',pages_total=1,pages_processed=1,
                       export_date='2027-02-23',invoice_number='S/NEW&01/27',destination_country='India',
                       consignee='NEW CONSIGNEE',consignee_address='ADDRESS ONE\nIEC & PAN : 000ABC\nGST : 00123',
                       carton_range='T01-T03, T05',package_count=4,gross_weight_kg=5.125,
                       items=[LineItem('UNSEEN FABRIC '+str(n),1.125,'YDS',hs_code='00123456',amount=0.23) for n in range(count)],
                       total_value=0 if count==0 else 0.23*count)

    def write(self,data,template=None):
        target=self.root/'invoice.xlsx'
        populate_template(template or ROOT/'templates/INDIA_template.xlsx',target,data)
        return target

    def test_zero_one_seven_and_forty_items_print_without_loss_or_footer_collision(self):
        for count in (0,1,7,40):
            with self.subTest(count=count):
                data=self.data(count);target=self.write(data)
                verify_workbook(target,data.to_dict(),self.pdf,'INDIA')
                book=load_workbook(target);s=book.active;extra=max(0,max(count,2)-6)
                self.assertEqual(s[f'H{42+extra}'].value,'TOTAL VALUE:')
                self.assertEqual(s[f'I{42+extra}'].value,data.total_value)
                self.assertIsInstance(s[f'I{42+extra}'].value,(int,float))
                self.assertIn(str(49+extra),str(s.print_area))
                for n in range(count):
                    self.assertEqual(s.cell(33+n,2).value,data.items[n].description)
                    self.assertEqual(s.cell(33+n,5).value,'00123456')
                    self.assertEqual(s.cell(33+n,5).data_type,'s')
                    self.assertEqual(s.cell(33+n,6).value,1.125)
                    self.assertEqual(s.cell(33+n,9).value,.23)
                self.assertEqual(s[f'A{47+extra}'].value.splitlines(),[
                    'THIS IS TRADE SAMPLES AND NOT FOR SALE.',
                    'THERE IS NO COMMERCIAL VALUE AND FOREIGN EXCHANGE INVOLVED.',
                    'VALUE DECLARED FOR CUSTOMS PURPOSE ONLY.'])
                occupied=set()
                for merge in s.merged_cells.ranges:
                    cells=set(merge.cells);self.assertFalse(occupied & cells);occupied.update(cells)
                book.close()

    def test_long_descriptions_addresses_marks_and_identifiers_remain_literal(self):
        data=self.data();data.items[0].description='=1+1\n'+'LONG COMPOSITION AND WIDTH '*25
        data.consignee='=UNKNOWN '+('COMPANY '*20)
        data.consignee_address='\n'.join('ADDRESS LINE '+str(n) for n in range(10))
        data.shipping_marks='LONG MARK '*6;data.carton_range='T01-T03, T05'
        book=load_workbook(self.write(data));s=book.active
        self.assertEqual(s['B33'].value,data.items[0].description);self.assertEqual(s['B33'].data_type,'s')
        self.assertGreater(s.row_dimensions[33].height,100)
        self.assertEqual(s['B11'].data_type,'s');self.assertGreater(s.row_dimensions[11].height,30)
        self.assertIn('ADDRESS LINE 9',s['B16'].value)
        self.assertGreater(s.row_dimensions[16].height,50)
        self.assertEqual((s['A33'].value,s['A34'].value),('T01-T03','T05'))
        self.assertEqual(s['H11'].value,data.shipping_marks);self.assertGreater(s.row_dimensions[11].height,30)
        # openpyxl misreads escaped ampersands followed by digits as a font-size code; verify the stored Excel text.
        with zipfile.ZipFile(self.root/'invoice.xlsx') as archive:
            self.assertIn('S/NEW&amp;&amp;01/27',archive.read('xl/worksheets/sheet1.xml').decode())
        self.assertIsNone(s.firstHeader.right.text)
        self.assertEqual(s.print_title_rows,'$30:$30');self.assertEqual(s.firstFooter.right.text,'Page &P of &N')
        book.close()

    def test_missing_values_clear_stale_cells_and_do_not_supply_destination_or_prices(self):
        stale=self.root/'INDIA_template.xlsx';book=load_workbook(ROOT/'templates/INDIA_template.xlsx')
        for coord in ('C9','B11','B12','B16','H11','H12','H13','D20','A17','A18','A22','A23','C25','B27','G27','B33','E33','F33','G33','I33','I42'):
            book.active[coord]='STALE'
        book.save(stale);book.close()
        data=FOCData(source_file=self.pdf.name,template_type='INDIA',pages_total=1,pages_processed=1)
        target=self.write(data,stale);verify_workbook(target,data.to_dict(),self.pdf,'INDIA')
        book=load_workbook(target);s=book.active
        for coord in ('C9','B11','B12','B16','H11','H12','H13','D20','C25','B27','G27','B33','E33','F33','G33','I33','I42'):
            self.assertIsNone(s[coord].value,coord)
        self.assertEqual(s['A17'].value,'SHIPMENT FROM:          TO:           BY:')
        self.assertEqual(s['A18'].value,'ON/ABT. ')
        self.assertEqual(s['A22'].value,'PORT OF DISCHARGE: ')
        self.assertEqual(s['A23'].value,'SHIPMENT TERMS: ')
        with zipfile.ZipFile(target) as archive:
            tree=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            self.assertFalse(any(not node.text for node in tree.findall(f'.//{namespace}r/{namespace}t')))
        self.assertFalse(any('REVIEW' in str(c.value) or 'VALIDATION' in str(c.value) for row in s for c in row))
        book.close()

    def test_shipment_section_uses_formatted_invoice_date_and_red_source_values(self):
        data=self.data();data.destination_city='NEW CITY, INDIA'
        data.port_of_discharge='SOURCE PORT';data.shipment_terms='FOB'
        target=self.write(data);verify_workbook(target,data.to_dict(),self.pdf,'INDIA')
        book=load_workbook(target,rich_text=True);s=book.active
        for coord,value in (('A17','NEW CITY, INDIA'),('A22','SOURCE PORT'),('A23','FOB')):
            self.assertTrue(all(v.text for v in s[coord].value if hasattr(v,'text')))
            runs=[v for v in s[coord].value if hasattr(v,'font') and v.text==value]
            self.assertEqual(len(runs),1)
            self.assertEqual(runs[0].font.color.rgb[-6:],'FF0000')
            self.assertGreaterEqual(s.row_dimensions[s[coord].row].height,20)
            self.assertNotEqual(s[coord].fill.fgColor.rgb,'FFFFFF00')
        self.assertEqual(str(s['A18'].value),'ON/ABT. FEB. 23, 2027')
        self.assertNotIn('HONG KONG',str(s['A17'].value))
        self.assertTrue(all(v.font.color.rgb[-6:]=='000000' for v in s['A18'].value if hasattr(v,'font')))
        self.assertEqual(s['C9'].value,datetime(2027,2,23))
        book.close()

    def test_typed_date_optional_port_terms_and_fractional_values(self):
        data=self.data();data.port_of_discharge='SOURCE PORT';data.shipment_terms='SOURCE TERMS'
        book=load_workbook(self.write(data));s=book.active
        self.assertEqual(s['C9'].value,datetime(2027,2,23));self.assertEqual(s['C9'].number_format,'dd-mmm-yyyy')
        self.assertEqual(s['G27'].value,5.125)
        self.assertEqual(s['A22'].value,'PORT OF DISCHARGE: SOURCE PORT')
        self.assertEqual(s['A23'].value,'SHIPMENT TERMS: SOURCE TERMS')
        self.assertEqual(s['B14'].font.color.rgb[-6:],'FF0000')
        self.assertEqual(s.page_setup.fitToWidth,1);self.assertEqual(s.page_setup.fitToHeight,0)
        book.close()

    def test_on_abt_tracks_the_invoice_date_without_a_fixed_sample_date(self):
        data=self.data()
        for value,expected in (('2026-10-14','ON/ABT. OCT. 14, 2026'),
                               ('2028-01-06','ON/ABT. JAN. 6, 2028'),(None,'ON/ABT. ')):
            data.export_date=value
            target=self.write(data);verify_workbook(target,data.to_dict(),self.pdf,'INDIA')
            book=load_workbook(target,rich_text=True)
            self.assertEqual(str(book.active['A18'].value),expected)
            book.close()

    def test_pipeline_accepts_india_growth_with_approved_rates(self):
        data=self.data(7);data.total_value=None
        data.port_of_discharge='SOURCE PORT';data.shipment_terms='FOB'
        for item in data.items:item.amount=None
        with patch('src.foc_extractor.pipeline.analyze_pdf',return_value=OCRResult(pages=['DESTINATION COUNTRY INDIA'])), \
             patch('src.foc_extractor.pipeline.parse_result',return_value=data):
            output=run(self.pdf,region='INDIA',output_dir=self.root/'output',local_model=None)
        verify_workbook(output,data.to_dict(),self.pdf,'INDIA')
        self.assertEqual(data.port_of_discharge,'CHENNAI, INDIA')
        self.assertEqual(data.shipment_terms,'CPT')
        book=load_workbook(output);self.assertEqual(book.active['I43'].value,2)
        self.assertTrue(all(item.unit_price==0.2 and item.amount==0.23 for item in data.items));book.close()

    def test_verifier_rejects_changed_total_text_date_and_missing_or_unprinted_declarations(self):
        for coordinate,value in (('I42',999),('C9','2027-02-23'),('A47',None)):
            with self.subTest(coordinate=coordinate):
                data=self.data();target=self.write(data);book=load_workbook(target)
                book.active[coordinate]=value;book.save(target);book.close()
                with self.assertRaises(ValueError):verify_workbook(target,data.to_dict(),self.pdf,'INDIA')
        data=self.data();target=self.write(data);book=load_workbook(target)
        book.active.print_area='A1:I43';book.save(target);book.close()
        with self.assertRaisesRegex(ValueError,'declarations'):verify_workbook(target,data.to_dict(),self.pdf,'INDIA')

    def test_integer_goods_quantities_do_not_print_a_trailing_decimal_point(self):
        data=self.data();data.items[0].quantity=378
        book=load_workbook(self.write(data))
        self.assertEqual(book.active['F33'].number_format,'#,##0');book.close()


if __name__=='__main__':unittest.main()
