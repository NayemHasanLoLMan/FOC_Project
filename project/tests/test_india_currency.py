from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.ocr_engine import OCRResult, OCRTable, OCRTableCell
from src.foc_extractor.parser import parse_result
from src.foc_extractor.pipeline import run
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.valuation import apply_valuation
from src.foc_extractor.workbook_verification import verify_workbook
from src.foc_extractor.pipeline_config import load_config
from src.foc_extractor.job_store import JobStore
from src.foc_extractor.queue_worker import QueueWorker


class IndiaCurrencyTests(unittest.TestCase):
    def ocr(self,content='',headers=('UNIT PRICE','AMOUNT'),values=('','')):
        rows=[['DESCRIPTION','QTY',*headers],['FABRIC','10 YDS',*values]]
        return OCRResult(content=content,pages=[content],total_pages=1,tables=[OCRTable(
            [OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)],2,4)])

    def parse(self,ocr):
        with patch('src.foc_extractor.parser._consignee_region',return_value='NEW INDIA FACTORY\nSOURCE ROAD'):
            return parse_result(ocr,Path('unseen.pdf'),region='INDIA')

    def test_labelled_currency_and_non_usd_never_get_usd_estimates(self):
        for text,expected in [('Currency: USD','USD'),('Currency EUR','EUR'),('Currency:\nINR','INR'),
                              ('Invoice currency = GBP','GBP'),('Currency: US Dollars','USD'),
                              ('Currency: US$','USD'),('Currency: XYZ','XYZ'),
                              ('Currency:\neur','EUR'),('Currency:\nEuro','EUR')]:
            with self.subTest(text=text):
                data=self.parse(self.ocr(text))
                self.assertEqual(data.currency,expected)
                apply_valuation(data)
                if expected!='USD':
                    self.assertIsNone(data.items[0].unit_price)
                    self.assertIsNone(data.total_value)

    def test_financial_headers_and_cells_detect_source_currency(self):
        for ocr in (self.ocr(headers=('UNIT PRICE (EUR)','AMOUNT (EUR)'),values=('3','30')),
                    self.ocr(values=('EUR 3.00','30 EUR')),
                    self.ocr('TOTAL VALUE: EUR 30.00'),self.ocr('AMOUNT: €30.00'),
                    self.ocr('UNIT PRICE: ₹3.00'),self.ocr('AMOUNT: HK$30.00'),
                    self.ocr(headers=('UNIT PRICE (XYZ)','AMOUNT'),values=('3','30'))):
            with self.subTest(content=ocr.content,table=ocr.tables[0]):
                data=self.parse(ocr)
                self.assertNotEqual(data.currency,'USD')
                apply_valuation(data)
                self.assertIsNone(data.items[0].unit_price)

    def test_conflicting_declarations_or_money_cells_raise(self):
        for ocr in (self.ocr('Currency USD\nTOTAL VALUE EUR 30'),
                    self.ocr('Currency: USD',headers=('UNIT PRICE (EUR)','AMOUNT')),
                    self.ocr(headers=('UNIT PRICE (USD)','AMOUNT'),values=('','EUR 30'))):
            with self.assertRaisesRegex(ValueError,'[Cc]onflicting.*currenc'):
                self.parse(ocr)

    def test_currency_table_cells_work_without_a_page_text_reading(self):
        for cells in ([OCRTableCell(0,0,'Currency'),OCRTableCell(0,1,'EUR')],
                      [OCRTableCell(0,0,'Currency'),OCRTableCell(1,0,'INR')],
                      [OCRTableCell(0,0,'Currency: GBP')]):
            ocr=self.ocr();ocr.tables.insert(0,OCRTable(cells,2,2))
            self.assertNotEqual(self.parse(ocr).currency,'USD')

    def test_unrecognized_explicit_currency_is_not_silently_usd(self):
        for text in ('Currency: unreadable','Currency: $','Currency: ¥','AMOUNT: ¥30','AMOUNT: Rs 30'):
            with self.subTest(text=text),self.assertRaisesRegex(ValueError,'[Cc]urrency'):
                self.parse(self.ocr(text))

    def test_absent_or_blank_currency_uses_documented_usd_default(self):
        for text in ('','Currency\nDocuments to accompany airwaybill','Country: INDIA\nEUR ROAD'):
            events=[]
            from src.foc_extractor.processing_evidence import observed
            @observed
            def extract():return self.parse(self.ocr(text))
            data=extract(observer=lambda event,**kwargs:events.append(event))
            self.assertEqual(data.currency,'USD')
            self.assertTrue(any(e.get('field')=='currency' and e.get('method')=='approved_default_usd' for e in events))
            apply_valuation(data)
            self.assertEqual(data.total_value,2)

    def test_explicit_usd_keeps_source_values_and_ocr_is_unchanged(self):
        ocr=self.ocr('Currency: USD',values=('USD 3.00','USD 29.75'));before=asdict(ocr)
        data=self.parse(ocr);apply_valuation(data)
        self.assertEqual(asdict(ocr),before)
        self.assertEqual((data.currency,data.items[0].unit_price,data.total_value),('USD',3,29.75))

    def test_pipeline_rejection_preserves_existing_excel_and_json(self):
        with tempfile.TemporaryDirectory() as directory:
            source=Path(directory)/'unseen.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            output=Path(directory)/'output';output.mkdir()
            target=output/'unseen.xlsx';target.write_bytes(b'previous invoice')
            data_path=output/'unseen.json';data_path.write_text('previous data')
            for content in ('Currency: EUR','Currency: USD\nTOTAL VALUE: EUR 30'):
                with patch('src.foc_extractor.pipeline.analyze_pdf',return_value=self.ocr(content)), \
                     patch('src.foc_extractor.parser._consignee_region',return_value='NEW INDIA FACTORY\nSOURCE ROAD'), \
                     patch('src.foc_extractor.pipeline.populate_template') as writer:
                    with self.assertRaises(ValueError):
                        run(source,region='INDIA',output_dir=output,save_json=True,local_model=None)
                writer.assert_not_called()
                self.assertEqual(target.read_bytes(),b'previous invoice')
                self.assertEqual(data_path.read_text(),'previous data')

    def test_direct_writer_and_verifier_reject_unsupported_currency(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);source=root/'unseen.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            target=root/'invoice.xlsx';target.write_bytes(b'previous invoice')
            data=FOCData(source_file=source.name,template_type='INDIA',currency='EUR',pages_total=1,pages_processed=1,
                         items=[LineItem('FABRIC',10,'YDS',amount=30)],total_value=30)
            with self.assertRaisesRegex(ValueError,'USD'):
                populate_template(ROOT/'templates/INDIA_template.xlsx',target,data)
            self.assertEqual(target.read_bytes(),b'previous invoice')
            with self.assertRaisesRegex(ValueError,'USD'):
                verify_workbook(target,data.to_dict(),source,'INDIA')

    def test_currency_audit_mismatch_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);source=root/'unseen.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            data=FOCData(source_file=source.name,template_type='INDIA',pages_total=1,pages_processed=1,
                         items=[LineItem('FABRIC',10,'YDS',amount=2)],total_value=2)
            target=root/'invoice.xlsx';populate_template(ROOT/'templates/INDIA_template.xlsx',target,data)
            book=load_workbook(target)
            for row in book['Extraction_Audit']:
                if row[0].value=='currency':row[1].value='EUR'
            with patch('src.foc_extractor.workbook_verification.load_workbook',return_value=book):
                with self.assertRaisesRegex(ValueError,'[Cc]urrency'):
                    verify_workbook(target,data.to_dict(),source,'INDIA')

    def test_queue_records_currency_failure_without_publishing_over_existing_invoice(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);settings=json.loads((ROOT/'pipeline_config.json').read_text());settings['local_model']=None
            config_path=root/'config.json';config_path.write_text(json.dumps(settings));config=load_config(config_path)
            store=JobStore(config.state_dir/'jobs.sqlite3')
            source=config.artifact_dir/'inputs/unseen.pdf';source.parent.mkdir(parents=True)
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            config.output_dir.mkdir(parents=True);existing=config.output_dir/'unseen.xlsx';existing.write_bytes(b'previous invoice')
            store.enqueue(config.incoming_dir/source.name,hashlib.sha256(source.read_bytes()).hexdigest(),source,region='INDIA')
            with patch('src.foc_extractor.pipeline.analyze_pdf',return_value=self.ocr('Currency: EUR')), \
                 patch('src.foc_extractor.parser._consignee_region',return_value='NEW INDIA FACTORY\nSOURCE ROAD'), \
                 QueueWorker(config,store) as worker:
                attempt_id=worker.run_once()
            attempt=store.get_attempt(attempt_id)
            self.assertEqual(attempt['state'],'failed')
            self.assertIn('source currency is EUR',attempt['error'])
            self.assertEqual(existing.read_bytes(),b'previous invoice')
            self.assertFalse(list((config.artifact_dir/'jobs').rglob('*.xlsx')))


if __name__=='__main__':unittest.main()
