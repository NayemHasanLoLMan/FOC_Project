from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pymupdf

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.parser import _extract_tin, _extract_import_cover_note, parse_result
from src.foc_extractor.ocr_engine import OCRResult
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.pipeline_config import load_config, initialize_folders
from src.foc_extractor.folder_intake import FolderScanner
from src.foc_extractor.job_store import JobStore


class BangladeshMetadataTests(unittest.TestCase):
    def test_tin_requires_a_label_and_preserves_suffixes_and_leading_zeros(self):
        for text,value in [('APPLICANT\'S TIN NO.472-620-435136 "FREIGHT PREPAID"','472-620-435136'),
                           ('TIN #140-200-1879/CIRCLE-50','140-200-1879/CIRCLE-50'),
                           ('TIN # 001-201-7957/Coy-20','001-201-7957/Coy-20'),
                           ('color# matte copper tin 3, 90pcs /color',None),
                           ('TIN #123-4567\nTIN #234-5678',None)]:
            self.assertEqual(_extract_tin(text),value)

    def test_cover_note_uses_only_the_supplied_text_with_bounded_continuation(self):
        note='Import open cover note # GDI/DEZ/07/2019/MAR/C/001352 (OPEN)'
        self.assertEqual(_extract_import_cover_note(note+'\nDated 03-JUL-2019\nCUSTOMS INVOICE NO.'),note+' Dated 03-JUL-2019')
        self.assertIsNone(_extract_import_cover_note('CONSIGNEE EGMCL\nBANGLADESH'))
        self.assertIsNone(_extract_import_cover_note('Import open cover note\nCUSTOMS INVOICE NO.'))

    def test_parser_does_not_seek_net_weight_or_invent_optional_metadata_for_bd(self):
        text='CONSIGNEE\nEPIC GARMENTS\nBANGLADESH\nNET WEIGHT 2 KG\nTIN #123-4567'
        data=parse_result(OCRResult(content=text,pages=[text]),Path('unseen.pdf'))
        self.assertIsNone(data.net_weight_kg)
        self.assertEqual(data.tin_number,'123-4567')
        self.assertIsNone(data.import_cover_note)

    def test_dynamic_red_fields_and_optional_footer_survive_table_growth(self):
        for count in (1,12):
            with self.subTest(count=count),tempfile.TemporaryDirectory() as directory:
                output=Path(directory)/'invoice.xlsx'
                note='Import open cover note # EXAMPLE/2027/C/1234 (OPEN) Dated 01-JAN-2027'
                data=FOCData(source_file='new.pdf',template_type='BD',tin_number='001-234-567890',
                    import_cover_note=note,consignee_address='Bangladesh\n(BIN:000322109-1402)',
                    carton_range='E01-E14',net_weight_kg=3,items=[LineItem('FABRIC',1,'YDS') for _ in range(count)])
                populate_template(ROOT/'templates/BD_template.xlsx',output,data)
                book=openpyxl.load_workbook(output,rich_text=True);ws=book.active
                self.assertEqual(ws['B13'].value,'BIN : 000322109-1402')
                self.assertEqual(ws['B13'].font.color.rgb[-6:],'FF0000')
                self.assertEqual(str(ws['H13'].value),'C/NO. E01-E14')
                self.assertEqual(ws['H13'].value[1].font.color.rgb[-6:],'FF0000')
                self.assertEqual(ws['A23'].value,note)
                self.assertEqual(ws['A23'].font.color.rgb[-6:],'FF0000')
                extra=max(0,count-6);self.assertEqual(ws.cell(41+extra,1).value,'TIN # 001-234-567890')
                self.assertEqual(ws.cell(41+extra,1).font.color.rgb[-6:],'FF0000')
                self.assertTrue(all(ws.cell(28,c).value is None for c in (6,7,8)))
                self.assertFalse(any('N.W.'==c.value for row in ws for c in row))
                self.assertNotIn('net_weight_kg',dict(book['Extraction_Audit'].iter_rows(min_row=2,max_col=2,values_only=True)))
                book.close()

    def test_default_country_folders_are_created_but_only_bd_is_watched(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'config.json';path.write_bytes((ROOT/'pipeline_config.json').read_bytes())
            config=load_config(path);initialize_folders(config)
            self.assertEqual(config.incoming_dir.name,'bangladesh')
            self.assertEqual(config.output_dir.name,'bangladesh')
            for base in ('incoming','output'):
                self.assertEqual({p.name for p in (path.parent/base).iterdir()},{'bangladesh','india','jordan'})

    def test_reserved_country_documents_are_not_queued_by_bd_scanner(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'config.json';path.write_bytes((ROOT/'pipeline_config.json').read_bytes())
            config=load_config(path);initialize_folders(config)
            with pymupdf.open() as pdf:
                pdf.new_page();content=pdf.tobytes()
            for country in ('bangladesh','india','jordan'):
                (config.incoming_dir.parent/country/'new.pdf').write_bytes(content)
            store=JobStore(config.state_dir/'jobs.sqlite3');scanner=FolderScanner(config,store)
            for when in (0,config.scan_interval_seconds):
                with patch('src.foc_extractor.folder_intake.time.time',return_value=when):
                    result=scanner.scan_once()
            self.assertEqual(store.counts()['queued'],1)
            self.assertEqual(Path(store.get_job(result['accepted_job_ids'][0])['source_path']).parent,config.incoming_dir)


if __name__=='__main__':unittest.main()
