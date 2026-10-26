import sys,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch,MagicMock
from PIL import Image,ImageDraw
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.foc_extractor.ocr_engine import OCRResult,OCRTable,OCRTableCell,OCRWord,_packing_footer
from src.foc_extractor.packages import packing_packages
from src.foc_extractor.weights import packing_weight,_weight_regions,_read_weight,_compact_weight_body,_correction_box


class PackingControlTests(unittest.TestCase):
    def compact(self):
        pages=['PACKING LIST\nNO OF PKGS DESCRIPTION SHIPPED QTY UNIT\nTotal: 1\nTotal Gross Weight (approx.): 0.5 Kgs.\nTotal Net Weight: 0.4 Kgs.',
               'PACKING LIST\nNO OF PKGS DESCRIPTION SHIPPED QTY UNIT\nTotal: 1\nTotal Gross Weight (approx.): 0.2 Kgs.']
        return OCRResult(pages=pages,total_pages=2,content='\n'.join(pages))

    def test_compact_local_totals_sum_once_and_do_not_become_shipment_total(self):
        o=self.compact()
        self.assertEqual(packing_packages(o)['total'],2)
        with tempfile.TemporaryDirectory() as directory:
            pdf=Path(directory)/'anonymous.pdf';pdf.write_bytes(b'source')
            result=packing_weight(o,pdf)
        self.assertEqual(result['total_kg'],.7)
        self.assertEqual(len(result['cells']),2)
        o.pages[0]=o.pages[0].replace('NO OF PKGS','')
        self.assertIsNone(_packing_footer(o,0,'package_count'))
        o.pages[0]+='\nTotal: 2'
        self.assertIsNone(_packing_footer(o,0,'package_count'))
        self.assertIsNone(packing_packages(o)['total'])

    def test_cover_without_weight_is_skipped_but_unknown_page_blocks_sum(self):
        o=self.compact();o.pages.insert(0,'SHIPPER\nCONSIGNEE\n2 PACKAGES');o.total_pages=3
        with tempfile.TemporaryDirectory() as directory:
            pdf=Path(directory)/'anonymous.pdf';pdf.write_bytes(b'source')
            self.assertEqual(packing_weight(o,pdf)['total_kg'],.7)
            o.pages.append('UNCLASSIFIED ATTACHMENT');o.total_pages=4
            self.assertIsNone(packing_weight(o,pdf)['total_kg'])
        o.pages[0]='SHIPPER\nCONSIGNEE'
        self.assertIsNone(packing_packages(o)['total'])

    def test_carton_identifier_does_not_replace_conflicting_explicit_counts(self):
        from test_packages import PackageTests
        o=OCRResult(pages=['PACKING LIST\nCARTON #14'],total_pages=1,
                    tables=[PackageTests().table(['1 pack','1 pack'])])
        result=packing_packages(o)
        self.assertEqual(result['total'],2)
        self.assertEqual(result['identifier_conflicts'],[{'page':1,'carton_ids':['14'],'explicit_count':2}])
        o.tables=[PackageTests().table(['PK'])]
        self.assertIsNone(packing_packages(o)['total'])

    def test_missing_packing_page_cannot_publish_partial_package_total(self):
        from test_packages import PackageTests
        o=OCRResult(pages=['PACKING LIST','PACKING LIST'],total_pages=2,
                    tables=[PackageTests().table(['1 pack'])])
        result=packing_packages(o)
        self.assertIsNone(result['total']);self.assertIn('pages: 2',result['reason'])

    def test_heading_declarations_cover_colour_continuations_above_quantities(self):
        image=Image.new('RGB',(600,270),'white');draw=ImageDraw.Draw(image)
        xs=[10,150,300,400,590];ys=[20,60,120,180,240]
        for x in xs:draw.line((x,20,x,240),fill='black',width=2)
        for y in ys:draw.line((10,y,590,y),fill='black',width=2)
        rows=[['DESCRIPTION','QTY','WEIGHT KGS','No of package / rolls'],
              ['100% COTTON','','0.2kg','1 pack'],['col. Red','3 YDS','',''],['col. Blue','4 YDS','','']]
        cells=[OCRTableCell(r,c,v,bbox=[xs[c],ys[r],xs[c+1],ys[r+1]])
               for r,row in enumerate(rows) for c,v in enumerate(row) if v]
        words=[OCRWord('3 YDS',.99,180,140,260,160),OCRWord('4 YDS',.99,180,200,260,220),
               OCRWord('0.2kg',.99,303,80,360,100),OCRWord('1 pack',.99,430,80,500,100)]
        o=OCRResult(pages=['PACKING LIST'],total_pages=1,words=words,tables=[OCRTable(cells,4,4,bbox=[10,20,590,240])])
        self.assertEqual(_weight_regions(image,o.tables[0],words),[[300,60,400,120]])
        reader=MagicMock();reader.predict.return_value=[{'rec_texts':['0.2kg'],'rec_scores':[.99]}]
        with tempfile.TemporaryDirectory() as directory:
            pdf=Path(directory)/'anonymous.pdf';pdf.write_bytes(b'source')
            with patch('src.foc_extractor.packages._cached_page_image',return_value=image):
                self.assertEqual(packing_packages(o,pdf=pdf)['total'],1)
            with patch('src.foc_extractor.weights._cached_page_image',return_value=image), \
                 patch('src.foc_extractor.weights._paddle_readers',return_value=(reader,None)):
                self.assertEqual(packing_weight(o,pdf)['total_kg'],.2)

    def test_subtotal_cells_are_not_added_to_their_detail_weights_or_packages(self):
        from test_packages import PackageTests
        image,o=PackageTests().physical(shared=False)
        t=o.tables[0];t.bbox[3]=230;t.row_count=4
        t.cells.extend([OCRTableCell(3,0,'Total:',bbox=[10,200,150,230]),
                        OCRTableCell(3,1,'5 PCS',bbox=[150,200,300,230]),
                        OCRTableCell(3,2,'3 KG',bbox=[300,200,400,230]),
                        OCRTableCell(3,3,'2 packs',bbox=[400,200,590,230])])
        o.words.extend([OCRWord('3 KG',.99,320,210,380,225),OCRWord('2 packs',.99,430,210,530,225)])
        reader=MagicMock();reader.predict.side_effect=[
            [{'rec_texts':['1kg'],'rec_scores':[.99]}],[{'rec_texts':['2kg'],'rec_scores':[.99]}]]
        with tempfile.TemporaryDirectory() as directory:
            pdf=Path(directory)/'anonymous.pdf';pdf.write_bytes(b'source')
            with patch('src.foc_extractor.packages._cached_page_image',return_value=image):
                self.assertEqual(packing_packages(o,pdf=pdf)['total'],2)
            with patch('src.foc_extractor.weights._cached_page_image',return_value=image), \
                 patch('src.foc_extractor.weights._paddle_readers',return_value=(reader,None)):
                self.assertEqual(packing_weight(o,pdf)['total_kg'],3)
                self.assertEqual(reader.predict.call_count,2)

    def test_positive_number_without_decimal_corroboration_is_not_a_verified_weight(self):
        from test_weights import WeightTests
        image,table=WeightTests().table(True)
        o=OCRResult(pages=['PACKING LIST'],total_pages=1,tables=[table])
        self.assertIsNone(_read_weight({'rec_texts':['B.KgS','35'],'rec_scores':[.99,.99]}))
        reader=MagicMock();reader.predict.return_value=[{'rec_texts':['B.KgS','35'],'rec_scores':[.99,.99]}]
        for numeric,literal,total in [(35,'3.5',None),(3.5,'3.5',3.5),(3.5,'35',None),(3.5,'3,5',None)]:
            with self.subTest(numeric=numeric,literal=literal),tempfile.TemporaryDirectory() as directory:
                pdf=Path(directory)/'anonymous.pdf';pdf.write_bytes(b'source')
                with patch('src.foc_extractor.weights._cached_page_image',return_value=image), \
                     patch('src.foc_extractor.weights._paddle_readers',return_value=(reader,None)), \
                     patch('src.foc_extractor.local_vision.transcribe_cell',side_effect=[
                         ({'weight_kg':numeric},'primary'),({'text':literal},'secondary')]):
                    result=packing_weight(o,pdf,local_model='primary')
                self.assertEqual(result['total_kg'],total)

    def test_compact_crop_keeps_decimal_pixels_without_blank_page_margins(self):
        image=Image.new('RGB',(170,1000),'white');draw=ImageDraw.Draw(image)
        draw.text((40,700),'3 5',fill='black');draw.point((48,711),fill='black')
        crop,box=_compact_weight_body(image)
        self.assertLess(crop.height,100)
        json.dumps(box)
        self.assertEqual(crop.getpixel((48-box[0],711-box[1])),(0,0,0))
        cell={'text':['B.KgS','35'],'ocr_boxes':[[10,20,60,40],[10,90,60,120]]}
        self.assertEqual(_correction_box(cell),[-6,74,76,136])
        cell['text']=['3.5','5']
        self.assertIsNone(_correction_box(cell))
        group={'text':['FAB','16kgs','25','8kgs','12'],
               'ocr_boxes':[[10,5,60,25],[10,20,60,40],[10,90,60,120],[10,220,60,240],[10,290,60,320]]}
        self.assertEqual(_correction_box(group,[0,0,100,150]),[-6,74,76,136])
        self.assertEqual(_correction_box(group,[0,200,100,350]),[-6,274,76,336])

    def test_inner_gross_label_does_not_block_pipeline_fallback(self):
        from src.foc_extractor.pipeline import run
        import openpyxl
        o=self.compact()
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);pdf=root/'anonymous.pdf';pdf.write_bytes(b'ocr-boundary')
            with patch('src.foc_extractor.pipeline.analyze_pdf',return_value=o), \
                 patch('src.foc_extractor.weights.packing_weight',return_value={'total_kg':.7,'cells':[],'reason':None}) as weights:
                out=run(pdf,output_dir=root/'out',local_model=None)
            book=openpyxl.load_workbook(out)
            self.assertEqual(book.active['G27'].value,.7);book.close();weights.assert_called_once()


if __name__=='__main__':unittest.main()
