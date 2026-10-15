import sys
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor.parser import (
    _extract_items_from_text, _extract_export_date, _extract_carton_range,
    _clean_material_description, _aggregate_items, _material_heading, _extract_gross_weight, parse_result,
)
from src.foc_extractor.ocr_engine import OCRResult, OCRTable, OCRTableCell
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.excel_writer import populate_template, _next_writable
from src.foc_extractor.pipeline import _resolve_template
from verify_extraction import read_invoice, compare
import openpyxl


class ParserTests(unittest.TestCase):
    def test_quantity_and_unit_are_distinct(self):
        for text,quantity,unit in [('COTTON 200 YDS',200,'YDS'),('LABEL 1,234.5 PCS',1234.5,'PCS'),('MOULD 5 SETS',5,'SETS')]:
            with self.subTest(text=text):
                item=_extract_items_from_text(text)[0]
                self.assertEqual((item.quantity,item.unit),(quantity,unit))

    def test_valid_and_ambiguous_dates(self):
        self.assertEqual(_extract_export_date('DATE OF EXPORT: 10/13/2027'),'2027-10-13')
        self.assertIsNone(_extract_export_date('DATE OF EXPORT: 10/09/2027'))
        self.assertIsNone(_extract_export_date('DATE OF EXPORT: 31/02/2027'))
        self.assertEqual(_extract_export_date('DATE OF EXPORTATION: 15-SEP,2027'),'2027-09-15')
        self.assertIsNone(_extract_export_date('PACKING DATE: 15-SEP,2027'))
        self.assertIsNone(_extract_export_date('no date'))
        self.assertEqual(_extract_export_date('PACKING DATE: 12-OCT-2027\nDATE OF EXPORT: 15-OCT-2027'),'2027-10-15')
        self.assertIsNone(_extract_export_date('DATE OF EXPORT: 15-OCT-2027\nDATE OF EXPORT: 16-OCT-2027'))

    def test_complete_numeric_tokens_and_composition(self):
        for value,expected in [('7,5',7.5),('1,234.5',1234.5),('1.234,5',1234.5)]:
            with self.subTest(value=value):
                self.assertEqual(_extract_items_from_text(f'COTTON {value} YDS')[0].quantity,expected)
                self.assertEqual(_extract_gross_weight(f'G.W. {value} KGS'),expected)
        self.assertEqual(_extract_items_from_text('COTTON -7 YDS'),[])
        self.assertIsNone(_material_heading('100% COTTON 5% UNKNOWN WOVEN'))
        self.assertIsNone(_material_heading('5% UNKNOWN 100% COTTON WOVEN'))
        self.assertEqual(_material_heading('STYLE SHIRT\n100% COTTON DENIM\nSHRINKAGE: 3%x3%', 'WOVEN'),'100% COTTON WOVEN FABRICS')
        self.assertEqual(_clean_material_description('9.5MM RIVET'),'9.5MM RIVET')

    def test_merged_description_and_source_rows_survive_aggregation(self):
        rows=[['PARTICULARS','QTY'],['100% COTTON WOVEN FABRICS',''],['','10 YDS'],['BLUE','2 YDS']]
        cells=[OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),2)]),Path('unseen.pdf'))
        self.assertEqual([(i.description,i.quantity) for i in data.items],[('100% COTTON WOVEN FABRICS',12)])
        self.assertEqual(data.items[0].source_page,1)
        self.assertEqual(len(data.items[0].source_refs),2)

    def test_late_header_and_explicit_item_column(self):
        rows=[['PACKING LIST'],[''],['NON MUTILATED'],[''],
              ['PARTICULARS','ITEM','QTY','WEIGHT KGS'],
              ['JACKET STYLE ABC\nSHANK BUTTON X 3 GROSS\nDIE MOULD X 2 SETS','SHANK BUTTON\nDIE MOULD','3 GROSS +\n2 SETS','4 KGS']]
        cells=[OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),4)]),Path('unseen.pdf'))
        self.assertEqual([(i.description,i.quantity,i.unit) for i in data.items],[('SHANK BUTTON',3,'GROSS'),('DIE MOULD',2,'SETS')])
        self.assertIsNone(data.gross_weight_kg)

    def test_multiline_garments_keep_quantities_with_compositions(self):
        rows=[['PARTICULARS','QTY'],['STYLE ABC SHIRT\n100% LINEN WOVEN SHIRT\n#50\nSTYLE DEF SHIRT\n60% VISCOSE 40% COTTON WOVEN SHIRT\n#30','3 PCS\n2 PCS']]
        cells=[OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),2)]),Path('unseen.pdf'))
        self.assertEqual([(i.description,i.quantity) for i in data.items],[('100% LINEN WOVEN SHIRT',3),('60% VISCOSE 40% COTTON WOVEN SHIRT',2)])

    def test_header_only_table_does_not_suppress_text(self):
        cells=[OCRTableCell(0,0,'PARTICULARS'),OCRTableCell(0,1,'QTY')]
        data=parse_result(OCRResult(tables=[OCRTable(cells,1,2)],pages=['COTTON 10 YDS']),Path('unseen.pdf'))
        self.assertEqual(data.items[0].quantity,10)

    def test_cartons_require_label(self):
        self.assertIsNone(_extract_carton_range('STYLE S3-S27'))
        self.assertIsNone(_extract_carton_range('C/NO. P09-P01'))
        self.assertEqual(_extract_carton_range('C/NO. E01-E13\nC/NO. B01'),'E01-E13, B01')
        self.assertEqual(_extract_carton_range('CINO. P1O P18'),'P10-P18')

    def test_explicit_packages_are_not_overwritten(self):
        data=parse_result(OCRResult(content='TOTAL 15 PACKAGES\nC/NO. E01-E01'),Path('arbitrary.pdf'))
        self.assertEqual(data.package_count,15)
        self.assertTrue(any('conflicts' in w for w in data.warnings))

    def test_empty_ocr_does_not_fabricate_data(self):
        data=parse_result(OCRResult(),Path('arbitrary.pdf'))
        self.assertEqual(data.items,[])
        for field in ['invoice_number','export_date','package_count','gross_weight_kg','total_value']:
            self.assertIsNone(getattr(data,field))

    def test_garments_and_multiword_materials_survive(self):
        self.assertEqual(_clean_material_description('MENS 100% LINEN WOVEN SHIRTS'),'MENS 100% LINEN WOVEN SHIRTS')
        desc='68% COTTON 28% RECYCLED POLYESTER 3% RAYON 1% SPANDEX WOVEN FABRICS'
        self.assertEqual(_clean_material_description(desc),desc)

    def test_aggregation_uses_real_unit(self):
        items=_aggregate_items([LineItem('COTTON',10,'YDS'),LineItem('COTTON',20,'YDS'),LineItem('COTTON',2,'PCS')])
        self.assertEqual([(i.quantity,i.unit) for i in items],[(30,'YDS'),(2,'PCS')])

    def test_table_columns_exclude_metadata_and_keep_first_row(self):
        cells=[OCRTableCell(0,c,t) for c,t in enumerate(['BUYER','PARTICULARS','QTY','WEIGHT KGS'])]
        cells += [OCRTableCell(1,c,t) for c,t in enumerate(['BUYER NAME','100% COTTON WOVEN FABRICS','25 YDS','2 KGS'])]
        result=OCRResult(tables=[OCRTable(cells,2,4)],content='PACKING LIST',pages=['PACKING LIST'])
        data=parse_result(result,Path('random.pdf'))
        self.assertEqual(len(data.items),1)
        self.assertEqual(data.items[0].quantity,25)
        self.assertNotIn('BUYER',data.items[0].description)

    def test_filename_is_not_silent_ground_truth(self):
        data=parse_result(OCRResult(),Path('27S9876-D.pdf'))
        self.assertEqual(data.invoice_number,'S/9876/27')
        self.assertTrue(any('filename' in w for w in data.warnings))

    def test_date_above_carton_label_is_not_package_count(self):
        data=parse_result(OCRResult(content='DATE 15 SEP 2027\nCARTONS'),Path('random.pdf'))
        self.assertIsNone(data.package_count)

    def test_color_rows_inherit_explicit_material_and_fabric_type(self):
        rows=[['PARTICULARS','QTY','WOVEN / KNIT'],
              ['100% NYLON','','WOVEN'],['BLUE','11 YDS',''],['RED','13 YDS',''],
              ['100% COTTON SHIRT','2 PCS','']]
        cells=[OCRTableCell(r,c,value) for r,row in enumerate(rows) for c,value in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),3)]),Path('unseen.pdf'))
        self.assertEqual([(i.description,i.quantity,i.unit) for i in data.items],
                         [('100% NYLON WOVEN FABRICS',24,'YDS'),('100% COTTON SHIRT',2,'PCS')])

    def test_narrow_ocr_repairs_are_flagged_and_ambiguous_digits_rejected(self):
        rows=[['PARTICULARS','QTY'],['LABEL','Ipc'],['CLOTH','9vds'],['OTHER','Zyds']]
        cells=[OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),2)]),Path('new.pdf'))
        self.assertEqual([(i.quantity,i.unit) for i in data.items],[(1,'PCS'),(9,'YDS'),(None,'YDS')])
        self.assertTrue(any('normalized OCR' in w for w in data.warnings))
        self.assertTrue(any('unresolved quantity' in w for w in data.warnings))

    def test_damaged_header_uses_neighboring_columns_not_row_totals(self):
        rows=[['PARIICULARS','90','WEIGHT KGS','No of package'],
              ['100% COTTON WOVEN FABRICS','25 vds','7 KGS','1 ROLL'],
              ['','', '','2027']]
        cells=[OCRTableCell(r,c,v) for r,row in enumerate(rows) for c,v in enumerate(row)]
        data=parse_result(OCRResult(tables=[OCRTable(cells,len(rows),4)]),Path('new.pdf'))
        self.assertEqual([(i.quantity,i.unit) for i in data.items],[(25,'YDS')])
        self.assertEqual(data.package_count,1)
        data=parse_result(OCRResult(content='C/NO. A01-A04',tables=[OCRTable(cells,len(rows),4)]),Path('new.pdf'))
        self.assertEqual(data.package_count,4)

    def test_shipping_marks_prioritize_consignee_over_incidental_mentions(self):
        data=parse_result(OCRResult(content='To: EGMCL-FOC\nPEARL BUTTONS'),Path('new.pdf'))
        self.assertEqual(data.consignee,'EPIC GARMENTS MANUFACTURING CO. LTD.')


class WriterTests(unittest.TestCase):
    def test_all_items_print_and_total_stays_below(self):
        data=FOCData(source_file='synthetic.pdf',template_type='BD',export_date='2027-01-23',
                     items=[LineItem(f'MATERIAL {i}',i+1,'PCS') for i in range(19)],warnings=['Review'])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            invoice=read_invoice(path)
            self.assertEqual(len(invoice['items']),19)
            self.assertTrue(invoice['audit_matches_invoice'])
            wb=openpyxl.load_workbook(path);ws=wb.active
            self.assertEqual(ws['F53'].value,'TOTAL VALUE:')
            self.assertIsNone(ws['G53'].value)
            self.assertIsNone(ws['A29'].value)
            self.assertIsInstance(ws['C9'].value,__import__('datetime').datetime)
            self.assertIn('56',str(ws.print_area))
            occupied=set()
            for merged in ws.merged_cells.ranges:
                cells=set(merged.cells)
                self.assertFalse(occupied & cells,'Overlapping merged cells can prevent Excel from opening the workbook')
                occupied.update(cells)
            wb.close()

    def test_missing_fields_are_blank_and_source_text_is_literal(self):
        data=FOCData(source_file='=1+1',template_type='BD',consignee='=1+1',
                     consignee_address='\n'.join(f'ADDRESS {i}' for i in range(8)),
                     items=[LineItem('=1+1',unit='PCS')],warnings=['REVIEW REQUIRED'])
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'invoice.xlsx'
            populate_template(ROOT/'templates/BD_template.xlsx',path,data)
            wb=openpyxl.load_workbook(path);ws=wb.active
            for coordinate in ['C9','H12','C25','B27','G27','G40','G33','A29']:
                self.assertIsNone(ws[coordinate].value,coordinate)
            for coordinate in ['B11','B33']:
                self.assertEqual(ws[coordinate].value,'=1+1')
                self.assertEqual(ws[coordinate].data_type,'s')
            self.assertIn('ADDRESS 7',ws['B17'].value)
            self.assertEqual(wb['Extracted_Items']['B2'].data_type,'s')
            self.assertFalse(any('REVIEW' in str(c.value) or 'DRAFT' in str(c.value) for row in ws for c in row))
            wb.close()

    def test_merged_label_fallback_avoids_nonanchor(self):
        ws=openpyxl.Workbook().active;ws.merge_cells('A1:B1');ws['A1']='LABEL'
        self.assertEqual(_next_writable(ws,ws['A1']).coordinate,'C1')

    def test_missing_explicit_template_fails(self):
        with self.assertRaises(FileNotFoundError):_resolve_template(Path('missing.xlsx'),'BD')

    def test_verifier_does_not_accept_nonempty_wrong_items(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);expected=root/'expected';output=root/'output'
            expected.mkdir();output.mkdir()
            for folder,qty in [(expected,10),(output,999)]:
                populate_template(ROOT/'templates/BD_template.xlsx',folder/'invoice.xlsx',
                                  FOCData(source_file='a.pdf',items=[LineItem('COTTON',qty,'YDS')],template_type='BD'))
            result=compare(output,expected)
            self.assertEqual(result['item_accuracy']['matched'],0)
            self.assertEqual(result['item_accuracy']['expected'],1)
            (output/'run_manifest.json').write_text(json.dumps([{'source':'a.pdf','status':'failed'}]))
            failed=compare(output,expected)
            self.assertEqual(failed['matched_files'],0)
            self.assertEqual(len(failed['incomplete_sources']),1)


class GeometryTests(unittest.TestCase):
    def test_deskew_accepts_installed_opencv_line_shape(self):
        from PIL import Image, ImageDraw
        from src.foc_extractor.ocr_engine import _deskew
        image=Image.new('RGB',(600,400),'white')
        draw=ImageDraw.Draw(image)
        for y in (50,150,250,350):draw.line((20,y,580,y),fill='black',width=2)
        self.assertEqual(_deskew(image).size,image.size)

    def test_grid_cells_use_word_coordinates(self):
        from PIL import Image, ImageDraw
        from src.foc_extractor.ocr_engine import _grid_tables, OCRWord
        image=Image.new('RGB',(800,500),'white');draw=ImageDraw.Draw(image)
        for x in (20,200,400,600,780):draw.line((x,40,x,460),fill='black',width=2)
        for y in (40,100,250,460):draw.line((20,y,780,y),fill='black',width=2)
        tables=_grid_tables(image,[OCRWord('PARTICULARS',.99,220,50,350,70),OCRWord('QTY',.99,420,50,460,70)],0)
        self.assertEqual(len(tables),1)
        self.assertEqual([(c.row_index,c.column_index,c.content) for c in tables[0].cells],[(0,1,'PARTICULARS'),(0,2,'QTY')])

    def test_grid_recovers_vertical_rules_interrupted_by_row_gutters(self):
        from PIL import Image, ImageDraw
        from src.foc_extractor.ocr_engine import _grid_tables, OCRWord
        image=Image.new('RGB',(800,500),'white');draw=ImageDraw.Draw(image)
        for x in (20,780):draw.line((x,40,x,460),fill='black',width=2)
        for x in (200,400,600):
            for top in range(40,460,60):draw.line((x,top,x,top+50),fill='black',width=2)
        for y in (40,100,160,220,280,340,400,460):draw.line((20,y,780,y),fill='black',width=2)
        tables=_grid_tables(image,[OCRWord('PARTICULARS',.99,220,50,350,70),OCRWord('QTY',.99,420,50,460,70)],0)
        self.assertEqual(tables[0].column_count,4)
        self.assertIsNotNone(tables[0].cells[0].bbox)


if __name__=='__main__':unittest.main()
