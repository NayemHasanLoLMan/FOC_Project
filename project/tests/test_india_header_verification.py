"""Publication checks the visible India header, not just its hidden audit sheet."""
import hashlib,json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor import pipeline,queue_worker
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.job_store import JobStore
from src.foc_extractor.models import FOCData,LineItem
from src.foc_extractor.ocr_engine import OCRResult
from src.foc_extractor.pipeline_config import load_config
from src.foc_extractor.workbook_verification import verify_workbook


class IndiaHeaderVerificationTests(unittest.TestCase):
    def setUp(self):
        temporary=tempfile.TemporaryDirectory();self.addCleanup(temporary.cleanup)
        self.root=Path(temporary.name).resolve();self.pdf=self.root/'unseen.pdf'
        with pymupdf.open() as pdf:
            pdf.new_page();pdf.save(self.pdf)
        self.data=FOCData(source_file=self.pdf.name,template_type='INDIA',pages_total=1,pages_processed=1,
                          invoice_number='S/12345/27',export_date='2027-03-12',consignee='UNSEEN FACTORY',
                          consignee_address='NEW ROAD\nIEC & PAN: 000ABCDE12\nGST: 001234567890123',
                          shipping_marks='NEW-FOC',destination_country='India',destination_city='MUMBAI, INDIA',
                          carton_range='T01-T03, T05',package_count=4,gross_weight_kg=5.125,
                          port_of_discharge='NEW PORT',shipment_terms='FOB',
                          items=[LineItem('ZIPPER',3,'PCS',amount=.3)],total_value=.3)
        self.target=self.root/'invoice.xlsx'

    def write(self,data=None):
        data=data or self.data
        populate_template(ROOT/'templates/INDIA_template.xlsx',self.target,data)
        return self.target

    def verify(self,data=None):
        verify_workbook(self.target,(data or self.data).to_dict(),self.pdf,'INDIA')

    def change(self,action):
        book=load_workbook(self.target)
        try:action(book.active);book.save(self.target)
        finally:book.close()

    def test_all_supported_headers_and_identifier_leading_zeros_pass(self):
        self.write();self.verify()

    def test_every_header_value_is_checked_independently_of_the_unchanged_audit(self):
        for coordinate in ('C9','C25','B11','B12','B13','B14','B15','B16','H11','H12','H13',
                           'D20','B27','G27','A17','A18','A22','A23','C27','H27'):
            with self.subTest(coordinate=coordinate):
                self.write();self.change(lambda s:s.__setitem__(coordinate,'WRONG'))
                with self.assertRaisesRegex(ValueError,'Printed India header'):self.verify()

    def test_missing_values_can_stay_blank_and_zero_values_are_preserved(self):
        blank=FOCData(source_file=self.pdf.name,template_type='INDIA',pages_total=1,pages_processed=1)
        self.write(blank);self.verify(blank)
        blank.package_count=0;blank.gross_weight_kg=0
        self.write(blank);self.verify(blank)
        self.change(lambda s:s.__setitem__('B27',True))
        with self.assertRaisesRegex(ValueError,'header B27'):self.verify(blank)

    def test_unsupported_missing_header_cannot_retain_a_stale_value(self):
        for coordinate in ('H11','H12','H13','C25','B12','A22','A23','B27','G27'):
            with self.subTest(coordinate=coordinate):
                blank=FOCData(source_file=self.pdf.name,template_type='INDIA',pages_total=1,pages_processed=1)
                self.write(blank);self.change(lambda s:s.__setitem__(coordinate,'STALE'))
                with self.assertRaisesRegex(ValueError,'Printed India header'):self.verify(blank)

    def test_long_address_folding_and_comma_split_remain_valid(self):
        for address in ('\n'.join('ADDRESS '+str(i) for i in range(8))+'\nGST: 000123456789012',
                        'A VERY LONG SOURCE STREET ADDRESS, DISTRICT AND CITY, STATE 000123, INDIA'):
            with self.subTest(address=address):
                self.data.consignee_address=address;self.write();self.verify()
                self.change(lambda s:s.__setitem__('B16' if '\n' in address else 'B12','MISSING PART'))
                with self.assertRaisesRegex(ValueError,'Printed India header'):self.verify()

    def test_literal_formula_like_identifiers_are_allowed_but_formulas_are_rejected(self):
        self.data.invoice_number='=1+1';self.write();self.verify()
        self.change(lambda s:s.__setitem__('C25','=1+1'))
        with self.assertRaisesRegex(ValueError,'literal readable format'):self.verify()

    def test_cartons_are_verified_in_the_goods_column_even_without_matching_goods_rows(self):
        for coordinate in ('A33','A34','A35'):
            with self.subTest(coordinate=coordinate):
                self.write();self.change(lambda s:s.__setitem__(coordinate,'WRONG'))
                with self.assertRaisesRegex(ValueError,'goods carton references'):self.verify()
        self.data.carton_range=', '.join(f'T{i:02}' for i in range(1,10))
        self.write();self.verify()

    def test_hidden_rows_columns_and_sheet_cannot_publish_populated_headers(self):
        actions=(lambda s:setattr(s.row_dimensions[11],'hidden',True),
                 lambda s:setattr(s.row_dimensions[18],'hidden',True),
                 lambda s:s.column_dimensions.group('B','D',hidden=True),
                 lambda s:setattr(s,'sheet_state','hidden'))
        for action in actions:
            with self.subTest(action=action):
                self.write();self.change(action)
                with self.assertRaisesRegex(ValueError,'must be visible'):self.verify()

    def test_header_and_carton_rows_must_be_in_the_print_area(self):
        self.write();self.change(lambda s:setattr(s,'print_area','A30:I49'))
        with self.assertRaisesRegex(ValueError,'print area'):self.verify()
        self.write();self.change(lambda s:setattr(s.row_dimensions[34],'hidden',True))
        with self.assertRaisesRegex(ValueError,'must be visible'):self.verify()

    def test_formats_that_hide_or_replace_header_values_are_rejected(self):
        for coordinate,formatting in (('G27',';;;'),('B27','"NOT A COUNT"'),('C25',';;;'),('C9','General')):
            with self.subTest(coordinate=coordinate):
                self.write();self.change(lambda s:setattr(s[coordinate],'number_format',formatting))
                with self.assertRaisesRegex(ValueError,'Printed India header'):self.verify()

    def test_direct_pipeline_failure_preserves_existing_excel_and_json(self):
        output=self.root/'output';output.mkdir()
        previous=output/'unseen.xlsx';previous.write_bytes(b'previous invoice')
        previous_json=output/'unseen.json';previous_json.write_bytes(b'previous data')
        def corrupt(*args,**kwargs):
            populate_template(*args,**kwargs)
            book=load_workbook(args[1]);book.active['C25']='WRONG';book.save(args[1]);book.close()
        with patch.object(pipeline,'analyze_pdf',return_value=OCRResult(pages=['DESTINATION COUNTRY INDIA'])), \
             patch.object(pipeline,'parse_result',return_value=self.data), \
             patch.object(pipeline,'populate_template',side_effect=corrupt):
            with self.assertRaisesRegex(ValueError,'header C25'):
                pipeline.run(self.pdf,region='INDIA',output_dir=output,local_model=None)
        self.assertEqual(previous.read_bytes(),b'previous invoice')
        self.assertEqual(previous_json.read_bytes(),b'previous data')

    def test_worker_failure_does_not_replace_a_previously_published_invoice(self):
        config_path=self.root/'pipeline_config.json'
        config_path.write_text((ROOT/'pipeline_config.json').read_text())
        config=load_config(config_path);store=JobStore(config.state_dir/'jobs.sqlite3')
        digest=hashlib.sha256(self.pdf.read_bytes()).hexdigest()
        first=store.enqueue(self.pdf,digest,self.pdf,region='INDIA')
        broken=False
        def convert(source,**kwargs):
            target=kwargs['output_dir']/'unseen.xlsx';populate_template(kwargs['template'],target,self.data)
            (kwargs['output_dir']/'unseen.json').write_text(json.dumps(self.data.to_dict()))
            if broken:
                book=load_workbook(target);book.active['B11']='WRONG';book.save(target);book.close()
            return target
        with patch.object(queue_worker,'run',side_effect=convert),queue_worker.QueueWorker(config,store) as worker:
            original_attempt=worker.run_once();published=store.get_attempt(original_attempt)
            self.assertEqual(published['state'],'completed',published['error'])
            previous=Path(published['output_path']);content=previous.read_bytes()
            source=self.root/'version2'/'unseen.pdf';source.parent.mkdir()
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((20,20),'NEW SOURCE');pdf.save(source)
            second=store.enqueue(self.pdf,hashlib.sha256(source.read_bytes()).hexdigest(),source,region='INDIA')
            broken=True;failed=store.get_attempt(worker.run_once())
        self.assertEqual(failed['state'],'failed');self.assertIn('header B11',failed['error'])
        self.assertEqual(previous.read_bytes(),content)
        self.assertEqual(store.get_job(first)['state'],'completed');self.assertEqual(store.get_job(second)['state'],'failed')
        self.assertIsNone(failed['output_path'])


if __name__=='__main__':unittest.main()
