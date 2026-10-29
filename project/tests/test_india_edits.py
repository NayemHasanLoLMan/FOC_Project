"""India edits reuse the real queue, publication, recovery and dashboard boundaries."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from uuid import uuid4

import pymupdf
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor import queue_worker
from src.foc_extractor.dashboard import create_app
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.invoice_edits import edited_data,validate_changes
from src.foc_extractor.job_store import JobStore
from src.foc_extractor.models import FOCData,LineItem
from src.foc_extractor.pipeline_config import load_config,attempt_metadata
from src.foc_extractor.queue_worker import QueueWorker
import test_invoice_edits as edit_fixtures


class IndiaEditTests(unittest.TestCase):
    submit=edit_fixtures.InvoiceEditTests.submit
    process_edit=edit_fixtures.InvoiceEditTests.process_edit

    def setUp(self):
        temporary=tempfile.TemporaryDirectory();self.addCleanup(temporary.cleanup)
        self.root=Path(temporary.name).resolve()
        config=self.root/'pipeline_config.json';config.write_text((ROOT/'pipeline_config.json').read_text())
        self.config=load_config(config);self.store=JobStore(self.config.state_dir/'jobs.sqlite3')
        self.source=self.config.artifact_dir/'inputs/india.pdf';self.source.parent.mkdir(parents=True)
        with pymupdf.open() as pdf:
            pdf.new_page();pdf.save(self.source)
        self.job=self.store.enqueue(self.config.incoming_dir/'india.pdf',hashlib.sha256(self.source.read_bytes()).hexdigest(),self.source,region='INDIA')
        self.client=create_app(self.config).test_client()
        self.headers={'Origin':'http://localhost','X-CSRF-Token':self.client.get('/api/controls').get_json()['csrf_token']}
        self.original=FOCData(source_file='india.pdf',template_type='INDIA',pages_total=1,pages_processed=1,
            invoice_number='S/12345/27',export_date='2027-02-03',destination_country='India',destination_city='NEW CITY, INDIA',
            consignee='ORIGINAL FACTORY',consignee_address='ORIGINAL ROAD\nIEC: 000ABCDE',
            carton_range='T01-T02',package_count=2,gross_weight_kg=4.5,total_value=2.50,
            items=[LineItem('FABRIC',2,'YDS',hs_code='00123456',unit_price=1,amount=2.23,source_page=1,source_refs=['page:1:row:1']),
                   LineItem('ZIPPER',2,'PCS',unit_price=.12,amount=.24,source_page=1,source_refs=['page:1:row:2'])])

    def complete(self):
        def convert(source,*,template,output_dir,observer,**kwargs):
            for kind,extra in [('vision_request',{}),('vision_response',{'answer':{'text':'2 PK'},'cached':True}),
                               ('vision_decision',{'candidate':'2 PK','accepted':True})]:
                observer({'kind':kind,'call_id':'original','field':'package_count','page':1,**extra})
            path=output_dir/'india.xlsx';populate_template(template,path,self.original)
            (output_dir/'india.json').write_text(json.dumps(self.original.to_dict()))
            return path
        with patch.object(queue_worker,'run',side_effect=convert),QueueWorker(self.config,self.store) as worker:
            attempt=worker.run_once()
        self.assertEqual(self.store.get_attempt(attempt)['state'],'completed',self.store.get_attempt(attempt)['error'])
        return attempt

    def rows(self):
        return [{'source_index':n,**{key:getattr(item,key) for key in ('description','quantity','unit','unit_price','hs_code')}}
                for n,item in enumerate(self.original.items)]

    def test_header_save_preserves_source_history_region_currency_and_old_download(self):
        old=self.complete();record=self.store.get_attempt(old)
        output=Path(record['output_path']);receipt=Path(record['receipt_path'])
        before=(output.read_bytes(),receipt.read_bytes(),self.source.read_bytes())
        changes={'consignee':'NEW FACTORY','consignee_address':'NEW ROAD\nGST: 000123',
                 'port_of_discharge':'USER PORT','shipment_terms':'FOB','gross_weight_kg':'5.125'}
        self.assertEqual(self.submit(old,changes).status_code,200)
        new=self.process_edit();detail=self.client.get(f'/api/jobs/{self.job}').get_json()
        self.assertNotEqual(output.read_bytes(),before[0])
        self.assertEqual((receipt.read_bytes(),self.source.read_bytes()),before[1:])
        self.assertEqual(Path(self.store.get_attempt(new)['output_path']),output)
        self.assertEqual(list(self.config.output_for('INDIA').glob('*.xlsx')),[output])
        self.assertEqual((detail['data']['template_type'],detail['data']['destination_country'],detail['data']['currency']),('INDIA','India','USD'))
        self.assertEqual(detail['data']['items'],self.original.to_dict()['items'])
        self.assertEqual(detail['data']['total_value'],2.50)
        book=load_workbook(self.store.get_attempt(new)['output_path']);s=book.active
        self.assertEqual(s['A22'].value,'PORT OF DISCHARGE: USER PORT');self.assertEqual(s['A23'].value,'SHIPMENT TERMS: FOB');book.close()
        for attempt in (old,new):
            self.assertEqual(self.client.get(f'/files/attempts/{attempt}/invoice').status_code,200)
            self.assertEqual(self.client.get(f'/invoices/attempts/{attempt}').status_code,200)

    def test_editor_offers_india_fields_and_hs_without_unprinted_bd_fields(self):
        self.complete();page=self.client.get(f'/jobs/{self.job}/edit')
        self.assertEqual(page.status_code,200)
        text=page.get_data(as_text=True)
        for value in ('Address / IEC / PAN / GST','name="port_of_discharge"','name="shipment_terms"','data-field="hs_code"','00123456','data-region="INDIA"'):
            self.assertIn(value,text)
        for key in ('tin_number','import_cover_note','destination_country','net_weight_kg','currency','template_type'):
            self.assertNotIn(f'name="{key}"',text)

    def test_goods_changes_and_new_rows_use_cents_leading_zeros_and_literal_text(self):
        old=self.complete();rows=self.rows()
        rows[0].update(description='=1+1',quantity='3',unit_price='0.335',hs_code='00000009')
        rows.append(dict(source_index=None,description='ADDED GOODS',quantity='1',unit='PCS',unit_price='.10',hs_code=''))
        self.assertEqual(self.submit(old,{'items':rows}).status_code,200)
        new=self.process_edit();data=self.client.get(f'/api/jobs/{self.job}').get_json()['data']
        self.assertEqual([i['amount'] for i in data['items']],[1.01,.24,.10]);self.assertEqual(data['total_value'],1.35)
        self.assertEqual(data['items'][0]['source_refs'],[]);self.assertIsNone(data['items'][0]['source_page'])
        self.assertEqual(data['items'][1]['source_refs'],['page:1:row:2'])
        book=load_workbook(self.store.get_attempt(new)['output_path']);s=book.active
        self.assertEqual((s['E33'].value,s['E33'].data_type),('00000009','s'))
        self.assertEqual((s['B33'].value,s['B33'].data_type),('=1+1','s'));self.assertEqual(s['I42'].value,1.35);book.close()

    def test_hs_or_description_only_preserves_pdf_amounts_and_declared_total(self):
        rows=self.rows();rows[0]['hs_code']='00000001';rows[1]['description']='SOURCE ZIPPER RENAMED'
        data=edited_data(self.original.to_dict(),{'items':rows})
        self.assertEqual([i.amount for i in data.items],[2.23,.24]);self.assertEqual(data.total_value,2.50)
        self.assertEqual(data.items[0].hs_code,'00000001');self.assertEqual(data.items[0].source_refs,[])
        self.assertEqual(self.original.items[0].hs_code,'00123456')
        original=self.original.to_dict();original['items'][0]['unit_price']=None
        rows[0]['unit_price']=''
        self.assertEqual(edited_data(original,{'items':rows}).items[0].amount,2.23)

    def test_blank_price_or_removed_goods_and_explicit_total_are_respected(self):
        rows=self.rows();rows[0].update(quantity=3,unit_price='',hs_code='')
        data=edited_data(self.original.to_dict(),{'items':rows})
        self.assertIsNone(data.items[0].amount);self.assertIsNone(data.total_value);self.assertIsNone(data.items[0].hs_code)
        self.assertEqual(edited_data(self.original.to_dict(),{'items':rows,'total_value':'9.99'}).total_value,9.99)
        self.assertIsNone(edited_data(self.original.to_dict(),{'items':[]}).total_value)

    def test_region_protected_values_and_invalid_hs_never_queue_work(self):
        old=self.complete()
        for changes in ({'destination_country':'Bangladesh'},{'template_type':'BD'},{'currency':'EUR'},
                        {'source_file':'other.pdf'},{'tin_number':'123'},{'import_cover_note':'note'},
                        {'shipment_terms':['FOB']},{'port_of_discharge':'x'*2001},
                        *({'items':[{**self.rows()[0],'hs_code':code}]} for code in (12,True,[], 'x'*21)),
                        {'items':[self.rows()[0],self.rows()[0]]}):
            with self.subTest(changes=repr(changes)[:80]):
                self.assertEqual(self.submit(old,changes).status_code,400)
        self.assertEqual(len(self.store.attempts(self.job)),1)
        with self.assertRaises(ValueError):validate_changes({'shipment_terms':'FOB'},'BD')
        with self.assertRaises(ValueError):validate_changes({'items':self.rows()},'BD')
        with self.assertRaises(ValueError):validate_changes({'consignee':'NEW'},'OTHER')

    def test_replay_stale_version_and_original_evidence_after_two_edits(self):
        old=self.complete();request_id=str(uuid4());changes={'shipment_terms':'FOB'}
        first=self.submit(old,changes,request_id);new=self.process_edit()
        self.assertEqual(self.submit(old,changes,request_id).get_json(),first.get_json())
        self.assertEqual(self.submit(old,changes).status_code,409)
        self.assertEqual(self.submit(old,{'shipment_terms':'CPT'},request_id).status_code,409)
        self.assertEqual(self.submit(new,{'port_of_discharge':'NEW PORT'}).status_code,200);self.process_edit()
        detail=self.client.get(f'/api/jobs/{self.job}').get_json()
        self.assertEqual(detail['evidence'][0]['reading'],'2 PK');self.assertTrue(detail['evidence'][0]['outcome'].startswith('Original reading:'))
        self.assertEqual(detail['data']['shipment_terms'],'FOB')

    def test_locked_save_and_restart_recovery_do_not_repeat_ocr(self):
        old=self.complete();self.assertEqual(self.submit(old,{'shipment_terms':'FOB'}).status_code,200)
        with patch.object(queue_worker,'run') as converter,QueueWorker(self.config,self.store) as worker:
            with patch.object(queue_worker,'_publish_copy',side_effect=PermissionError('Locked')):
                new=worker.run_once()
            self.assertEqual(self.store.get_attempt(new)['state'],'waiting_to_save')
        with patch.object(queue_worker,'run') as converter2,QueueWorker(self.config,self.store) as worker:
            worker.run_once()
        converter.assert_not_called();converter2.assert_not_called()
        self.assertEqual(self.store.get_attempt(new)['state'],'completed')
        self.assertEqual(self.submit(new,{'shipment_terms':'CPT'}).status_code,200)
        with QueueWorker(self.config,self.store) as worker:
            metadata=attempt_metadata(self.config,'INDIA');metadata['worker_lock']=str(worker.lock_path)
            interrupted=self.store.claim_next(worker.owner,metadata)
        final=self.process_edit()
        self.assertEqual(self.store.get_attempt(interrupted['id'])['state'],'interrupted')
        self.assertEqual(self.store.get_attempt(final)['metadata']['manual_edit']['base_attempt_id'],new)

    def test_failed_verification_preserves_the_published_original(self):
        old=self.complete();previous=Path(self.store.get_attempt(old)['output_path']);content=previous.read_bytes()
        self.assertEqual(self.submit(old,{'shipment_terms':'FOB'}).status_code,200)
        def corrupt(*args,**kwargs):
            populate_template(*args,**kwargs)
            book=load_workbook(args[1]);book.active['E33']='WRONG';book.save(args[1]);book.close()
        with patch.object(queue_worker,'run') as converter,patch.object(queue_worker,'populate_template',side_effect=corrupt),QueueWorker(self.config,self.store) as worker:
            failed=self.store.get_attempt(worker.run_once())
        converter.assert_not_called();self.assertEqual(failed['state'],'failed');self.assertIsNone(failed['output_path'])
        self.assertEqual(previous.read_bytes(),content)
        self.assertEqual(self.client.get(f'/files/attempts/{old}/invoice').status_code,200)
