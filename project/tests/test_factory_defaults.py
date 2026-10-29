"""Approved regional details must not change PDF goods or inactive regions."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.foc_extractor import factory_profiles
from src.foc_extractor.factory_profiles import identify_factory, apply_defaults, apply_jordan_origin
from src.foc_extractor.models import FOCData, LineItem
from src.foc_extractor.ocr_engine import OCRResult
from src.foc_extractor.processing_evidence import observed
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.invoice_edits import edited_data
from src.foc_extractor.workbook_verification import verify_workbook

FIELDS = ('consignee','consignee_address','destination_country','tin_number','import_cover_note','shipment_text','remarks')


class FactoryDefaultsTests(unittest.TestCase):
    def setUp(self):
        self.master = json.loads(factory_profiles.MASTER_FILE.read_text(encoding='utf-8'))

    def identity(self,text='TO: PGCL',region='BD'):
        return identify_factory(OCRResult(content=text,pages=[text]),region)

    def data(self,count=1,region='BD'):
        return FOCData(source_file='anonymous.pdf',template_type=region,pages_processed=1,pages_total=1,
                       consignee='PDF FACTORY',consignee_address='PDF ROAD\nBIN : 123456789-0000',
                       tin_number='123-456-789',import_cover_note='PDF COVER NOTE',destination_country='Bangladesh',
                       destination_city='PDF CITY',export_date='2026-10-20',invoice_number='S/9012/26',
                       shipping_marks='SOURCE-FABRIC',carton_range='A01-A02',package_count=2,gross_weight_kg=1.25,
                       total_value=10,items=[LineItem('SOURCE ITEM '+str(i),2,'PCS',unit_price=.25,amount=.5) for i in range(count)])

    def test_rose_replaces_only_supplied_fields_and_preserves_source_values(self):
        data=self.data();data.shipment_text='PDF SHIPMENT';before=data.to_dict()
        identity=self.identity('TO: THE ROSE DRESSES LTD')
        self.assertEqual(identity['factory_id'],'ROSE')
        apply_defaults(data,identity)
        profile=self.master['factories']['ROSE']
        for field in ('consignee','consignee_address','remarks'):
            self.assertEqual(getattr(data,field),profile[field])
        for field,value in before.items():
            if field not in ('consignee','consignee_address','remarks','warnings'):
                self.assertEqual(data.to_dict()[field],value,field)

    def test_all_five_bd_profiles_replace_only_owned_fields_and_audit_originals(self):
        for code in ('EGMCL','EG2','GTL','CIPL','PGCL'):
            with self.subTest(code=code):
                data = self.data();before = data.to_dict();events = []
                data.warnings = ['consignee_address not established from source; review required.','goods missing']
                identity = self.identity('TO: '+code)
                observed(apply_defaults)(data,identity,observer=lambda event,**kw:events.append(event))
                profile = self.master['factories'][code]
                for field in FIELDS:
                    self.assertEqual(getattr(data,field),profile[field])
                    event = next(e for e in events if e.get('field')==field)
                    self.assertEqual(event['original_value'],before[field])
                    self.assertEqual(event['profile_sha256'],identity['profile_sha256'])
                for field,value in before.items():
                    if field not in FIELDS+('warnings',):self.assertEqual(getattr(data,field) if field!='items' else data.to_dict()['items'],value)
                self.assertIn('goods missing',data.warnings)
                self.assertFalse(any(w.startswith('consignee_address not established') for w in data.warnings))

    def test_unknown_conflicts_and_other_regions_are_unchanged(self):
        for text,region in [('TO: THE UNLISTED DRESSES LTD','BD'),('TO: PGCL\nCONSIGNEE: CIPL','BD'),
                            ('TO: TGPL','BD'),('TO: UNKNOWN INDIA FACTORY','INDIA'),('TO: UNKNOWN JORDAN FACTORY','JORDAN')]:
            with self.subTest(text=text,region=region):
                data = self.data(region=region);before = deepcopy(data)
                apply_defaults(data,self.identity(text,region))
                self.assertEqual(data,before)

    def test_changed_master_is_rejected_and_inactive_defaults_do_nothing(self):
        data = self.data();identity = self.identity();before = deepcopy(data)
        with patch.object(Path,'read_bytes',return_value=b'{}'):
            with self.assertRaisesRegex(ValueError,'changed'):apply_defaults(data,identity)
        self.assertEqual(data,before)
        master = deepcopy(self.master);master['defaults_regions'] = []
        content = json.dumps(master).encode()
        with patch.object(Path,'read_bytes',return_value=content):
            apply_defaults(data,{**identity,'profile_sha256':hashlib.sha256(content).hexdigest()})
        self.assertEqual(data,before)

    def test_explicit_no_need_omits_source_values_and_wording(self):
        data = self.data();identity = self.identity()
        for field in ('tin_number','import_cover_note','shipment_text','remarks'):
            self.master['factories']['PGCL'][field] = None
        content = json.dumps(self.master).encode()
        with patch.object(Path,'read_bytes',return_value=content):
            apply_defaults(data,{**identity,'profile_sha256':hashlib.sha256(content).hexdigest()})
        self.assertIsNone(data.tin_number);self.assertIsNone(data.import_cover_note)
        self.assertEqual(data.shipment_text,'');self.assertEqual(data.remarks,'')
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/data.source_file;workbook = root/'omitted.xlsx'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            populate_template(ROOT/'templates/BD_template.xlsx',workbook,data)
            verify_workbook(workbook,data.to_dict(),source,'BD')
            book = load_workbook(workbook)
            for coordinate in ('A18','A23','A41','A42','A43'):self.assertIsNone(book.active[coordinate].value)
            book.close()

    def test_invalid_printed_values_fail_before_changing_data(self):
        for value in ('',123,'x'*2001):
            with self.subTest(value=str(value)[:20]):
                data = self.data();before = deepcopy(data);master = deepcopy(self.master)
                master['factories']['PGCL']['consignee_address'] = value
                content = json.dumps(master).encode();identity = self.identity()
                with patch.object(Path,'read_bytes',return_value=content):
                    with self.assertRaises(ValueError):apply_defaults(data,{**identity,'profile_sha256':hashlib.sha256(content).hexdigest()})
                self.assertEqual(data,before)

    def test_bd_workbooks_print_master_details_and_verify_after_growth(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/'anonymous.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            for code in ('EGMCL','EG2','GTL','CIPL','PGCL'):
                for count in (1,9):
                    with self.subTest(code=code,count=count):
                        data = self.data(count);apply_defaults(data,self.identity('TO: '+code))
                        path = root/(code+str(count)+'.xlsx')
                        populate_template(ROOT/'templates/BD_template.xlsx',path,data)
                        verify_workbook(path,data.to_dict(),source,'BD')
                        book = load_workbook(path);sheet = book.active;total = 40+max(0,count-6)
                        self.assertEqual(sheet['A18'].value,data.shipment_text)
                        self.assertEqual(sheet.cell(total+2,1).value,'SAMPLE OF NO COMMERCIAL VALUE')
                        self.assertEqual(sheet.cell(total+3,1).value,'INVOICE ONLY FOR CUSTOMS PURPOSE')
                        bins = [sheet.cell(r,2) for r in range(12,18) if str(sheet.cell(r,2).value or '').startswith('BIN :')]
                        self.assertEqual(len(bins),1);self.assertEqual(bins[0].font.color.rgb,'FFFF0000')
                        if data.tin_number:self.assertEqual(sheet.cell(total+1,1).font.color.rgb,'FFFF0000')
                        if data.import_cover_note:self.assertEqual(sheet['A23'].font.color.rgb,'FFFF0000')
                        self.assertFalse(any('review required' in str(c.value).lower() for row in sheet for c in row))
                        self.assertEqual((sheet['F28'].value,sheet['G28'].value,sheet['H28'].value),(None,None,None))
                        # Corrupting a printed profile field must prevent publication.
                        sheet['B11'] = 'WRONG FACTORY';book.save(path);book.close()
                        with self.assertRaisesRegex(ValueError,'Bangladesh header'):verify_workbook(path,data.to_dict(),source,'BD')

    def test_manual_edit_preserves_overrides_and_stored_wording(self):
        data = self.data();apply_defaults(data,self.identity())
        edited = edited_data(data.to_dict(),{'consignee':'MANUAL FACTORY','consignee_address':'MANUAL ADDRESS','tin_number':'999-111-222'})
        self.assertEqual(edited.consignee,'MANUAL FACTORY');self.assertEqual(edited.consignee_address,'MANUAL ADDRESS')
        self.assertEqual(edited.tin_number,'999-111-222')
        self.assertEqual(edited.shipment_text,data.shipment_text);self.assertEqual(edited.remarks,data.remarks)
        self.assertEqual(edited_data({'source_file':'old.pdf','template_type':'BD','items':[]},{'consignee':'MANUAL'}).shipment_text,None)

    def test_india_profile_replaces_only_approved_fields_and_preserves_pdf_identifiers_in_evidence(self):
        data = self.data(region='INDIA');data.items[0].origin_country = 'VIETNAM'
        data.port_of_discharge = 'PDF PORT';data.shipment_terms = 'FOB'
        data.consignee_address = 'PDF ROAD\nIEC & PAN : OLD1234567\nGST Number : OLD TAX'
        before = data.to_dict();events = []
        identity = self.identity('CONSIGNEE: TGPL','INDIA')
        observed(apply_defaults)(data,identity,observer=lambda event,**kw:events.append(event))
        for field in FIELDS:
            self.assertEqual(getattr(data,field),self.master['factories']['INDIA'][field])
            event = next(e for e in events if e.get('field')==field)
            self.assertEqual(event['original_value'],before[field]);self.assertEqual(event['factory_id'],'INDIA')
        for field,value in before.items():
            if field not in FIELDS+('warnings',):self.assertEqual(data.to_dict()[field],value)
        self.assertIn('AAJCT7390N',data.consignee_address)
        self.assertIn('21AAJCT7390N1ZS',data.consignee_address)
        self.assertIsNone(data.tin_number);self.assertIsNone(data.import_cover_note)

    def test_india_defaults_require_explicit_enablement_and_a_country_consistent_match(self):
        data = self.data(region='INDIA');before = deepcopy(data)
        master = deepcopy(self.master);master['defaults_regions'] = ['BD']
        content = json.dumps(master).encode();identity = self.identity('TO: TGPL','INDIA')
        with patch.object(Path,'read_bytes',return_value=content):
            apply_defaults(data,{**identity,'profile_sha256':hashlib.sha256(content).hexdigest()})
        self.assertEqual(data,before)
        for text in ('TO: UNKNOWN FACTORY','TO: PGCL','TO: TGPL\nCONSIGNEE: PGCL'):
            apply_defaults(data,self.identity(text,'INDIA'));self.assertEqual(data,before)

    def test_india_master_layout_verifies_after_growth_and_rejects_changed_fields(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/'anonymous.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            for count in (1,9,40):
                data = self.data(count,region='INDIA');apply_defaults(data,self.identity('TO: TGPL','INDIA'))
                path = root/('india'+str(count)+'.xlsx')
                populate_template(ROOT/'templates/INDIA_template.xlsx',path,data)
                verify_workbook(path,data.to_dict(),source,'INDIA')
                book = load_workbook(path,rich_text=True);sheet = book.active;extra = max(0,count-6)
                self.assertEqual(str(sheet['A17'].value),data.shipment_text)
                runs = [run for run in sheet['A17'].value if hasattr(run,'font') and run.text=='CHENNAI']
                self.assertEqual(len(runs),1);self.assertEqual(runs[0].font.color.rgb,'FFFF0000')
                for coordinate in ('B14','B15'):self.assertEqual(sheet[coordinate].font.color.rgb,'FFFF0000')
                self.assertEqual(sheet[f'A{47+extra}'].value,data.remarks)
                self.assertEqual(str(sheet['A18'].value),'ON/ABT. OCT. 20, 2026')
                self.assertEqual(sheet['A22'].value,'PORT OF DISCHARGE: ')
                self.assertEqual(sheet['A23'].value,'SHIPMENT TERMS: ')
                book.close()
                for coordinate in ('B14','B15','A17',f'A{47+extra}'):
                    populate_template(ROOT/'templates/INDIA_template.xlsx',path,data)
                    book = load_workbook(path);book.active[coordinate] = 'WRONG';book.save(path);book.close()
                    with self.assertRaisesRegex(ValueError,'India header'):verify_workbook(path,data.to_dict(),source,'INDIA')

    def test_india_explicit_omissions_and_saved_wording_survive_manual_edits(self):
        data = self.data(region='INDIA');apply_defaults(data,self.identity('TO: TGPL','INDIA'))
        edited = edited_data(data.to_dict(),{'consignee':'MANUAL FACTORY','consignee_address':'MANUAL ADDRESS'})
        self.assertEqual(edited.consignee,'MANUAL FACTORY');self.assertEqual(edited.consignee_address,'MANUAL ADDRESS')
        self.assertEqual(edited.shipment_text,data.shipment_text);self.assertEqual(edited.remarks,data.remarks)
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/data.source_file;path = root/'india.xlsx'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            populate_template(ROOT/'templates/INDIA_template.xlsx',path,edited)
            verify_workbook(path,edited.to_dict(),source,'INDIA')
            for wording in ('','CUSTOM SAVED SHIPMENT'):
                edited.shipment_text = wording;edited.remarks = ''
                populate_template(ROOT/'templates/INDIA_template.xlsx',path,edited)
                verify_workbook(path,edited.to_dict(),source,'INDIA')
                book = load_workbook(path)
                self.assertEqual(book.active['A17'].value,wording or None);self.assertIsNone(book.active['A47'].value)
                book.close()

    def test_pipeline_applies_india_master_without_changing_goods_or_source_pricing(self):
        from src.foc_extractor import pipeline
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/'anonymous.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            data = self.data(region='INDIA');data.destination_country = 'India';data.invoice_number = 'S/9012/26'
            before = deepcopy(data.items);events = []
            ocr = OCRResult(content='CONSIGNEE: TGPL\nPACKING LIST S/9012/26',pages=['CONSIGNEE: TGPL'],total_pages=1)
            with patch.object(pipeline,'analyze_pdf',return_value=ocr),patch.object(pipeline,'parse_result',return_value=data):
                path = pipeline.run(source,region='INDIA',output_dir=root/'out',local_model=None,save_json=True,
                                    observer=lambda event,**kw:events.append(event))
            verify_workbook(path,data.to_dict(),source,'INDIA')
            self.assertEqual(data.items,before);self.assertEqual(data.total_value,10)
            self.assertEqual(data.consignee_address,self.master['factories']['INDIA']['consignee_address'])
            saved = json.loads((root/'out/anonymous.json').read_text(encoding='utf-8'))
            self.assertEqual(saved['shipment_text'],self.master['factories']['INDIA']['shipment_text'])
            self.assertTrue(any(e.get('method')=='approved_factory_master' and e.get('original_value')=='PDF FACTORY' for e in events))

    def test_jordan_header_defaults_and_regional_origin_keep_original_readings(self):
        data = self.data(4,region='JORDAN');data.currency = 'EUR'
        for index,origin in enumerate((None,'VIETNAM','BANGLADESH','CHINA')):
            data.items[index].origin_country = origin;data.items[index].source_page = index+1
            data.items[index].source_refs = [str(index+1)+':table:row']
        before = data.to_dict();events = [];identity = self.identity('TO: EJU1','JORDAN')
        observed(apply_defaults)(data,identity,observer=lambda event,**kw:events.append(event))
        for field in FIELDS:
            expected = self.master['factories']['JORDAN'][field]
            if field=='shipment_text':expected = ''
            self.assertEqual(getattr(data,field),expected)
        observed(apply_jordan_origin)(data,identity,observer=lambda event,**kw:events.append(event))
        for index,item in enumerate(data.items):
            self.assertEqual(item.origin_country,'CHINA')
            self.assertEqual({k:v for k,v in data.to_dict()['items'][index].items() if k!='origin_country'},
                             {k:v for k,v in before['items'][index].items() if k!='origin_country'})
        for field,value in before.items():
            if field not in FIELDS+('items','warnings'):self.assertEqual(data.to_dict()[field],value)
        origins = [event for event in events if event.get('policy')=='jordan_always_china']
        self.assertEqual([event['original_value'] for event in origins],[None,'VIETNAM','BANGLADESH','CHINA'])
        self.assertTrue(all(event['method']=='approved_regional_rule' and event['source_refs'] for event in origins))

    def test_jordan_origin_is_regional_even_when_factory_is_unknown(self):
        data = self.data(region='JORDAN');data.destination_country = 'Jordan';data.items[0].origin_country = 'VIETNAM'
        identity = self.identity('TO: UNLISTED JORDAN CUSTOMER','JORDAN');before = deepcopy(data)
        apply_defaults(data,identity);self.assertEqual(data,before)
        events = []
        observed(apply_jordan_origin)(data,identity,observer=lambda event,**kw:events.append(event))
        self.assertEqual(data.items[0].origin_country,'CHINA')
        self.assertEqual((data.consignee,data.consignee_address),(before.consignee,before.consignee_address))
        self.assertIsNone(next(event for event in events if event.get('field')=='origin_country')['factory_id'])
        for region,text in (('BD','TO: PGCL'),('INDIA','TO: TGPL')):
            data = self.data(region=region);data.items[0].origin_country = 'VIETNAM';before = deepcopy(data)
            apply_jordan_origin(data,self.identity(text,region));self.assertEqual(data,before)

    def test_jordan_origin_requires_enabled_approved_policy_and_stable_master(self):
        data = self.data(region='JORDAN');identity = self.identity('TO: EJU1','JORDAN');before = deepcopy(data)
        master = deepcopy(self.master);master['defaults_regions'] = ['BD','INDIA']
        content = json.dumps(master).encode()
        with patch.object(Path,'read_bytes',return_value=content):
            apply_jordan_origin(data,{**identity,'profile_sha256':hashlib.sha256(content).hexdigest()})
        self.assertEqual(data,before)
        for change in ('origin','policy','checksum'):
            master = deepcopy(self.master)
            if change=='origin':master['factories']['JORDAN']['origin_country'] = 'VIETNAM'
            if change=='policy':master['policies']['jordan_origin'] = 'fill_missing'
            content = json.dumps(master).encode()
            updated = identity if change=='checksum' else {**identity,'profile_sha256':hashlib.sha256(content).hexdigest()}
            with patch.object(Path,'read_bytes',return_value=content):
                with self.assertRaises(ValueError):apply_jordan_origin(data,updated)
            self.assertEqual(data,before)

    def test_jordan_master_layout_and_origins_verify_after_growth(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/'anonymous.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            for count in (0,1,9,40):
                data = self.data(count,region='JORDAN');identity = self.identity('TO: EJU1','JORDAN')
                apply_defaults(data,identity);apply_jordan_origin(data,identity)
                path = root/('jordan'+str(count)+'.xlsx')
                populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
                verify_workbook(path,data.to_dict(),source,'JORDAN')
                book = load_workbook(path);sheet = book.active;extra = max(0,count-6)
                self.assertEqual(sheet[f'A{37+extra}'].value,'SAMPLE OF NO COMMERCIAL VALUE')
                self.assertEqual(sheet[f'A{38+extra}'].value,'INVOICE ONLY FOR CUSTOMS PURPOSE')
                self.assertEqual([sheet.cell(27+i,2).value for i in range(count)],
                                 [item.description for item in data.items])
                self.assertFalse(any('SHIPMENT FROM' in str(cell.value) for row in sheet for cell in row))
                book.close()
                for coordinate in ('C12',f'A{37+extra}'):
                    populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
                    book = load_workbook(path);book.active[coordinate] = 'WRONG';book.save(path);book.close()
                    with self.assertRaises(ValueError):verify_workbook(path,data.to_dict(),source,'JORDAN')
                if count:
                    populate_template(ROOT/'templates/JORDAN_template.xlsx',path,data)
                    book = load_workbook(path);book.active['B27'] = 'WRONG ORIGIN';book.save(path);book.close()
                    with self.assertRaisesRegex(ValueError,'goods'):verify_workbook(path,data.to_dict(),source,'JORDAN')

    def test_jordan_omitted_declarations_and_manual_versions_keep_stored_values(self):
        data = self.data(region='JORDAN');identity = self.identity('TO: EJU1','JORDAN')
        apply_defaults(data,identity);apply_jordan_origin(data,identity)
        edited = edited_data(data.to_dict(),{'consignee':'MANUAL CUSTOMER','items':[{
            'source_index':0,'description':data.items[0].description,'quantity':data.items[0].quantity,
            'unit':data.items[0].unit,'unit_price':data.items[0].unit_price,'weight_kg':.75,'origin_country':'VIETNAM'}]})
        self.assertEqual(edited.items[0].origin_country,'VIETNAM');self.assertEqual(edited.consignee,'MANUAL CUSTOMER')
        self.assertEqual(edited.remarks,data.remarks);self.assertEqual(edited.total_value,data.total_value)
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/data.source_file;path = root/'jordan.xlsx'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            for remarks in (data.remarks,'','CUSTOM SAVED DECLARATION\nSECOND LINE\nTHIRD LINE'):
                edited.remarks = remarks
                populate_template(ROOT/'templates/JORDAN_template.xlsx',path,edited)
                verify_workbook(path,edited.to_dict(),source,'JORDAN')

    def test_pipeline_applies_jordan_origin_after_source_enrichment(self):
        from src.foc_extractor import pipeline
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary);source = root/'anonymous.pdf'
            with pymupdf.open() as pdf:pdf.new_page();pdf.save(source)
            data = self.data(region='JORDAN');data.destination_country = 'Jordan';events = []
            ocr = OCRResult(content='SHIPPER\nS/9012/26\nCONSIGNEE: EPIC DESIGNERS LTD JORDAN',
                            pages=['SHIPPER\nS/9012/26\nCONSIGNEE: EPIC DESIGNERS LTD JORDAN'],total_pages=1)
            def packing(data,*args):
                data.items[0].origin_country = 'VIETNAM';data.items[0].weight_kg = .75
                return {'packages':{'total':2,'basis':'cover','reason':None,'entries':[]},'cells':[],'gross_basis':'shipment cover'}
            with patch.object(pipeline,'analyze_pdf',return_value=ocr),patch.object(pipeline,'parse_result',return_value=data), \
                 patch('src.foc_extractor.jordan.reconcile_packing',side_effect=packing), \
                 patch('src.foc_extractor.jordan.resolve_date',return_value=data.export_date):
                path = pipeline.run(source,region='JORDAN',output_dir=root/'out',local_model=None,save_json=True,
                                    observer=lambda event,**kw:events.append(event))
            verify_workbook(path,data.to_dict(),source,'JORDAN')
            self.assertEqual((data.items[0].origin_country,data.items[0].weight_kg,data.total_value),('CHINA',.75,10))
            event = next(e for e in events if e.get('policy')=='jordan_always_china')
            self.assertEqual(event['original_value'],'VIETNAM')
            saved = json.loads((root/'out/anonymous.json').read_text(encoding='utf-8'))
            self.assertEqual(saved['items'][0]['origin_country'],'CHINA')


if __name__=='__main__':unittest.main()
