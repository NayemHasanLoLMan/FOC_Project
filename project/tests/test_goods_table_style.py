from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.foc_extractor.excel_writer import populate_template
from src.foc_extractor.models import FOCData, LineItem


class GoodsTableStyleTests(unittest.TestCase):
    def test_logical_borders_follow_growth_without_changing_values_or_red_marks(self):
        with tempfile.TemporaryDirectory() as temporary:
            for region in ('BD', 'INDIA'):
                for count in (0, 1, 9, 40):
                    with self.subTest(region=region, count=count):
                        data = FOCData(source_file='unseen.pdf', template_type=region,
                                       carton_range='A01-A02', total_value=1.25,
                                       items=[LineItem('LONG FABRIC DESCRIPTION '*5, 1.25, 'YDS',
                                                       hs_code='00123456', amount=.25) for _ in range(count)])
                        path = Path(temporary)/'invoice.xlsx'
                        populate_template(ROOT/f'templates/{region}_template.xlsx', path, data)
                        book = load_workbook(path); sheet = book.active
                        last = max(33, 32+count)
                        for row in (30, 33, last):
                            self.assertEqual(sheet.cell(row, 1).border.left.style, 'thin')
                            self.assertEqual(sheet.cell(row, 9).border.right.style, 'thin')
                            self.assertEqual(sheet.cell(row, 2).border.bottom.style, 'thin')
                            self.assertIsNone(sheet.cell(row, 3).border.left.style)
                        self.assertEqual(sheet['B30'].fill.fgColor.rgb, 'FFE9EDF2')
                        if last < 38:
                            self.assertIsNone(getattr(sheet.cell(last+1, 1).border.left, 'style', None))
                        self.assertEqual(sheet['A33'].value, 'A01-A02')
                        self.assertEqual(sheet['A33'].font.color.rgb[-6:], 'FF0000')
                        if count:
                            self.assertEqual(sheet.cell(32+count, 2).value, data.items[-1].description)
                            self.assertEqual(sheet.cell(32+count, 7 if region == 'BD' else 6).value, 1.25)
                        total = 40 if region == 'BD' else 42
                        self.assertEqual(sheet.cell(total+max(0,count-6), 7 if region == 'BD' else 9).value, 1.25)
                        self.assertEqual(sheet.print_title_rows, '$30:$30')
                        book.close()
