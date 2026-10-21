"""Read the invoice worksheet for a local, non-editing HTML preview."""
from datetime import date, datetime
import io
import re

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils.cell import range_boundaries, get_column_letter


def _color(value, default='222222'):
    rgb = getattr(value,'rgb',None)
    return rgb[-6:] if isinstance(rgb,str) and re.fullmatch(r'[0-9a-fA-F]{6,8}',rgb) else default


def _font(font):
    size = min(36,max(8,float(getattr(font,'sz',None) or 11)))
    return f'color:#{_color(font.color)};font-size:{size}pt;font-weight:{700 if font.b else 400};font-style:{"italic" if font.i else "normal"}'


def _value(value, number_format):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.strftime('%d %b %Y')
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        if 'USD' in number_format:return f'USD {value:,.2f}'
        if number_format=='#,##0':return f'{value:,.0f}'
        if number_format=='#,##0.####':return f'{value:,.4f}'.rstrip('0').rstrip('.')
    return str(value)


def invoice_sheet(content):
    if len(content)>20*1024*1024:
        raise ValueError('Workbook is too large for a browser preview')
    book=load_workbook(io.BytesIO(content),data_only=True,rich_text=True)
    try:
        sheet=book.worksheets[0]
        area=sheet.print_area
        if area and ',' in area:raise ValueError('Multiple print areas are not supported')
        bounds=range_boundaries(area.rsplit('!',1)[-1]) if area else (1,1,sheet.max_column,sheet.max_row)
        left,top,right,bottom=bounds
        if right-left>31 or bottom-top>999:raise ValueError('Worksheet is too large for a browser preview')
        styles={}
        def style(cell):
            key='cell-'+str(cell.style_id)
            if key not in styles:
                align=cell.alignment.horizontal
                align=align if align in ('left','center','right','justify') else 'right' if isinstance(cell.value,(int,float)) else 'left'
                rules=[_font(cell.font),'text-align:'+align,'vertical-align:'+('middle' if cell.alignment.vertical=='center' else 'bottom' if cell.alignment.vertical=='bottom' else 'top')]
                if cell.fill.patternType=='solid':rules.append('background:#'+_color(cell.fill.fgColor,'FFFFFF'))
                for edge in ('top','bottom','left','right'):
                    border=getattr(cell.border,edge)
                    if border and border.style:
                        pattern={'dashed':'dashed','dotted':'dotted','double':'double'}.get(border.style,'solid')
                        width=2 if border.style in ('medium','thick','double') else 1
                        rules.append(f'border-{edge}:{width}px {pattern} #{_color(border.color)}')
                styles[key]=';'.join(rules)
            return key
        merged={};skip=set()
        for region in sheet.merged_cells.ranges:
            if region.min_col<left or region.max_col>right or region.min_row<top or region.max_row>bottom:continue
            merged[(region.min_row,region.min_col)]=(region.max_row-region.min_row+1,region.max_col-region.min_col+1)
            skip.update((r,c) for r in range(region.min_row,region.max_row+1) for c in range(region.min_col,region.max_col+1) if (r,c)!=(region.min_row,region.min_col))
        rows=[]
        for r in range(top,bottom+1):
            cells=[]
            for c in range(left,right+1):
                if (r,c) in skip:continue
                cell=sheet.cell(r,c);parts=[]
                if isinstance(cell.value,CellRichText):
                    for part in cell.value:
                        if isinstance(part,str):parts.append({'text':part,'class':''})
                        else:
                            key='run-'+str(len(styles));styles[key]=_font(part.font)
                            parts.append({'text':part.text,'class':key})
                else:parts=[{'text':_value(cell.value,cell.number_format),'class':''}]
                rowspan,colspan=merged.get((r,c),(1,1))
                cells.append({'class':style(cell),'parts':parts,'rowspan':rowspan,'colspan':colspan})
            key='row-'+str(r);styles[key]='height:'+str(min(500,max(1,sheet.row_dimensions[r].height or 15)))+'pt'
            if sheet.row_dimensions[r].hidden:styles[key]+=';display:none'
            rows.append({'class':key,'cells':cells})
        columns=[]
        for c in range(left,right+1):
            key='col-'+str(c);width=sheet.column_dimensions[get_column_letter(c)].width or 13
            styles[key]='width:'+str(min(700,max(4,width*7+5)))+'px';columns.append(key)
        return {'rows':rows,'columns':columns,'styles':styles,'sheet_name':sheet.title}
    finally:book.close()
