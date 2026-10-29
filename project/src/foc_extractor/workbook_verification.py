"""Check the printed invoice independently of its extracted-items sheet."""
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.print_settings import PrintArea
from datetime import date, datetime
import re
import pymupdf

from .regions import REGIONS, region_code, INDIA_DECLARATIONS


def verify_workbook(workbook, data, source, region='BD'):
    region_code(region)
    if region=='INDIA' and data.get('currency')!='USD':
        raise ValueError('India invoice publication requires USD source/default currency')
    if data.get('template_type') not in (region, None if region == 'BD' else region):
        raise ValueError('Converter result does not match the queued invoice region')
    if data.get('source_file') != source.name or not isinstance(data.get('items'), list):
        raise ValueError('Converter result does not identify this input snapshot')
    with pymupdf.open(source) as pdf:
        pages = len(pdf)
    if data.get('pages_processed') != pages or data.get('pages_total') != pages:
        raise ValueError('Refusing to publish an invoice with incomplete source-page coverage')
    book = load_workbook(workbook)
    try:
        audit = dict(book['Extraction_Audit'].iter_rows(min_row=2, max_col=2, values_only=True))
        if region=='INDIA' and audit.get('currency')!='USD':
            raise ValueError('India workbook currency does not match USD conversion result')
        if audit.get('source_file') != source.name or audit.get('pages_processed') != str(pages):
            raise ValueError('Workbook audit does not match the conversion result')
        if audit.get('template_type') not in ({region, None, 'None'} if region == 'BD' else {region}):
            raise ValueError('Workbook audit does not match the invoice region')
        sheet = book.active
        shipment = data.get('shipment_text') if data.get('shipment_text') is not None else 'SHIPMENT TO: ' + (
            data.get('destination_city') or data.get('destination_country') or '')
        if region == 'BD' and (sheet['A18'].value != (shipment or None) or (shipment and sheet['A18'].data_type != 's')):
            raise ValueError('Printed Bangladesh shipment destination does not match the conversion result')
        if region=='INDIA' and (sheet!=book.worksheets[0] or sheet.sheet_state!='visible'):
            raise ValueError('Printed India invoice sheet must be visible and active')
        layout = REGIONS[region]
        totals = [row for row in range(layout['start'], sheet.max_row + 1)
                  if sheet.cell(row, layout['total_column']).value == 'TOTAL VALUE:']
        if len(totals) != 1:
            raise ValueError('Printed invoice must have one identifiable total below its goods')
        total_row = totals[0]
        if region=='BD' and data.get('shipment_text') is not None:
            from .excel_writer import _split_address_lines
            address = [re.sub(r'^[\[(]?\s*BIN\s*[:#]?\s*([0-9 -]+?)\s*[)\]]?$',
                              lambda match:'BIN : '+match[1].strip(),line,flags=re.I)
                       for line in _split_address_lines(data.get('consignee_address'))]
            if len(address)>6:address = address[:5]+['\n'.join(address[5:])]
            headers = {'B11':data.get('consignee'),'E21':(data.get('destination_country') or '').upper(),
                       'A23':data.get('import_cover_note'),
                       f'A{total_row+1}':'TIN # '+data['tin_number'] if data.get('tin_number') else None}
            headers.update({f'B{12+i}':address[i] if i<len(address) else None for i in range(6)})
            if data.get('remarks') is not None:
                lines = data['remarks'].splitlines()
                headers.update({f'A{total_row+2}':lines[0] if lines else None,
                                f'A{total_row+3}':'\n'.join(lines[1:]) or None})
            for coordinate,value in headers.items():
                cell = sheet[coordinate]
                if cell.value!=(value or None) or (value and cell.data_type!='s'):
                    raise ValueError('Printed Bangladesh header '+coordinate+' does not match the conversion result')
        columns = [2, layout['quantity'], layout['unit']]
        keys = ['description', 'quantity', 'unit']
        combined_jordan = region=='JORDAN' and 'G24:H24' in sheet.merged_cells
        if combined_jordan:
            columns = [2, layout['quantity']]
            keys = ['description', 'quantity']
        priced_bd = region == 'BD' and sheet['G30'].value == 'UNIT PRICE (USD)'
        if priced_bd:
            columns = [2, 5, 6, 7, 9]
            keys += ['unit_price', 'amount']
        if region == 'INDIA':
            columns += [5, 9]; keys += ['hs_code', 'amount']
        elif region == 'JORDAN':
            columns += [6, 9]; keys += ['weight_kg', 'amount']
        actual = [tuple(sheet.cell(row, column).value for column in columns)
                  for row in range(layout['start'], total_row) if sheet.cell(row, 2).value is not None]
        expected = [tuple(item.get(key) or None if key in ('unit', 'hs_code') else item.get(key) for key in keys)
                    for item in data['items']]
        if region=='JORDAN':
            if combined_jordan:
                quantities = []
                for index,item in enumerate(data['items']):
                    row = layout['start']+index
                    if f'G{row}:H{row}' not in sheet.merged_cells:
                        raise ValueError('Printed Jordan quantity and unit must share one column')
                    quantity = f'{item["quantity"]:,.4f}'.rstrip('0').rstrip('.') if item.get('quantity') is not None else ''
                    quantities.append(' '.join(part for part in (quantity,item.get('unit')) if part) or None)
                expected = [(row[0],quantity,*row[2:]) for row,quantity in zip(expected,quantities)]
            # Accept retained pre-upgrade invoices as well as the uniform-origin footer.
            origins = {item.get('origin_country') for item in data['items']}
            common = next(iter(origins)) if len(origins)==1 and None not in origins and '' not in origins else None
            footer = sheet.cell(total_row+1,1).value
            if common and footer == 'COUNTRY OF ORIGIN: '+common:
                pass
            else:
                expected=[(row[0]+('\nCOUNTRY OF ORIGIN: '+item['origin_country'] if item.get('origin_country') else ''),*row[1:])
                          for row,item in zip(expected,data['items'])]
                if footer and str(footer).startswith('COUNTRY OF ORIGIN:'):
                    raise ValueError('Printed Jordan origin does not match the conversion result')
        if actual != expected:
            raise ValueError('Printed goods do not match the conversion result')
        if region != 'BD':
            areas = [range_boundaries(str(area)) for area in PrintArea.from_string(sheet.print_area).ranges] if sheet.print_area else []
            if not any(left <= 1 and right >= (10 if region=='JORDAN' else 9) and top <= 1 and bottom >= total_row
                       for left, top, right, bottom in areas):
                raise ValueError('Regional print area does not include all goods and the total')
            if region == 'INDIA':
                if sheet.cell(total_row, 9).value != data.get('total_value'):
                    raise ValueError('Printed India total does not match the conversion result')
                expected_date = date.fromisoformat(data['export_date'][:10]) if data.get('export_date') else None
                headers = {
                    'C9':expected_date, 'C25':data.get('invoice_number'), 'B11':data.get('consignee'),
                    'H11':data.get('shipping_marks'), 'H12':data.get('destination_city'),
                    'H13':'C/NO. '+data['carton_range'] if data.get('carton_range') else None,
                    'D20':data['destination_country'].upper() if data.get('destination_country') else None,
                    'B27':data.get('package_count'), 'G27':data.get('gross_weight_kg'),
                    'A17':data.get('shipment_text') if data.get('shipment_text') is not None else
                          'SHIPMENT FROM:          TO: '+(data.get('destination_city') or '')+'          BY:',
                    'A18':'ON/ABT. '+(f'{expected_date:%b}. {expected_date.day}, {expected_date.year}'.upper() if expected_date else ''),
                    'A22':'PORT OF DISCHARGE: '+(data.get('port_of_discharge') or ''),
                    'A23':'SHIPMENT TERMS: '+(data.get('shipment_terms') or ''),
                    'C27':'PKG(S) ONLY.', 'H27':'KGS',
                }
                if data.get('remarks') is not None:headers[f'A{total_row+5}'] = data['remarks']
                address = [line.strip() for line in (data.get('consignee_address') or '').splitlines() if line.strip()]
                if len(address)==1 and len(address[0])>60:
                    address = [part.strip() for part in address[0].split(',') if part.strip()]
                if len(address)>5:
                    address = address[:4]+['\n'.join(address[4:])]
                headers.update({f'B{12+index}':address[index] if index<len(address) else None for index in range(5)})
                cartons = [part.strip() for part in (data.get('carton_range') or '').split(',') if part.strip()]
                actual_cartons = [sheet.cell(row,1).value for row in range(layout['start'],total_row)]
                if actual_cartons != cartons+[None]*max(0,len(actual_cartons)-len(cartons)):
                    raise ValueError('Printed India goods carton references do not match the conversion result')
                headers.update({f'A{layout["start"]+index}':value for index,value in enumerate(cartons)})
                for coordinate,expected_value in headers.items():
                    cell = sheet[coordinate]
                    value = cell.value.date() if isinstance(cell.value,datetime) else cell.value
                    if value != (None if expected_value=='' else expected_value) or isinstance(value,bool):
                        raise ValueError(f'Printed India header {coordinate} does not match the conversion result')
                    if value is None:
                        continue
                    formats = {'General','@'} if isinstance(value,str) else {'General','0.00############'}
                    if cell.column==1 and cell.row>=layout['start'] and isinstance(value,str):
                        formats.add('d-mmm')  # The master uses a numeric date format; literal carton text prints unchanged.
                    if coordinate=='C9':
                        formats = {'dd-mmm-yyyy'}
                    if (cell.number_format not in formats or (isinstance(value,str) and cell.data_type!='s')):
                        raise ValueError(f'Printed India header {coordinate} must retain its literal readable format')
                    dimension = sheet.row_dimensions[cell.row]
                    if (sheet.sheet_state!='visible' or dimension.hidden or dimension.height==0
                            or any(d.hidden or d.width==0 for d in sheet.column_dimensions.values()
                                   if d.min<=cell.column<=d.max)
                            or not any(left<=cell.column<=right and top<=cell.row<=bottom
                                       for left,top,right,bottom in areas)):
                        raise ValueError(f'Printed India header {coordinate} must be visible inside the print area')
                declarations = data['remarks'].splitlines() if data.get('remarks') is not None else INDIA_DECLARATIONS
                for declaration in declarations:
                    cells = [cell for row in sheet for cell in row if declaration in str(cell.value or '').splitlines()]
                    if len(cells) != 1 or cells[0].row <= total_row or not any(
                            left <= cells[0].column <= right and top <= cells[0].row <= bottom
                            for left, top, right, bottom in areas):
                        raise ValueError('India declarations must appear once below the total inside the print area')
            if region == 'JORDAN':
                if sheet != book.worksheets[0] or sheet.sheet_state != 'visible':
                    raise ValueError('Printed Jordan invoice sheet must be visible and active')
                if sheet.cell(total_row, 9).value != data.get('total_value'):
                    raise ValueError('Printed Jordan total does not match the conversion result')
                headers = {
                    'C9': date.fromisoformat(data['export_date'][:10]) if data.get('export_date') else None,
                    'C11': data.get('consignee'), 'I9': data.get('shipping_marks'),
                    'I10': data.get('destination_city'), 'I11': data.get('carton_range'),
                    'D16': data['destination_country'].upper() if data.get('destination_country') else None,
                    'C18': data.get('invoice_number'), 'B21': data.get('package_count'),
                    'F21': data.get('gross_weight_kg'), 'I24': f'VALUE ({data.get("currency", "USD")})',
                    'A16': 'COUNTRY OF ULTIMATE DESTINATION:', 'A18': 'CUSTOMS INVOICE NO.',
                    'B24': 'DESCRIPTION OF GOODS', 'F24': 'KG', 'G24': 'QUANTITY', 'H24': None if combined_jordan else 'UNIT',
                }
                address = [line.strip() for line in (data.get('consignee_address') or '').splitlines() if line.strip()]
                if len(address) == 1 and len(address[0]) > 60:
                    address = [part.strip() for part in address[0].split(',') if part.strip()]
                if len(address) > 4:
                    address = address[:3] + ['\n'.join(address[3:])]
                headers.update({f'C{12+i}': address[i] if i < len(address) else None for i in range(4)})
                if data.get('remarks') is not None:
                    lines = data['remarks'].splitlines()
                    headers.update({f'A{total_row+2}':lines[0] if lines else None,
                                    f'A{total_row+3}':'\n'.join(lines[1:]) or None})
                for coordinate, expected_value in headers.items():
                    cell = sheet[coordinate]
                    value = cell.value.date() if isinstance(cell.value, datetime) else cell.value
                    if value != (None if expected_value == '' else expected_value) or isinstance(value, bool):
                        raise ValueError(f'Printed Jordan header {coordinate} does not match the conversion result')
                cartons = [part.strip() for part in (data.get('carton_range') or '').split(',') if part.strip()]
                actual_cartons = [sheet.cell(row, 1).value for row in range(layout['start'], total_row)]
                if actual_cartons != cartons + [None] * max(0, len(actual_cartons) - len(cartons)):
                    raise ValueError('Printed Jordan goods carton references do not match the conversion result')
                checked = [sheet[coordinate] for coordinate in headers]
                checked += [sheet.cell(total_row, 9)]
                checked += [sheet.cell(row, col) for row in range(layout['start'], total_row)
                            for col in (1, 2, 6, 7, 8, 9)]
                for cell in checked:
                    if cell.value is None:
                        continue
                    dimension = sheet.row_dimensions[cell.row]
                    if (dimension.hidden or dimension.height == 0
                            or any(d.hidden or d.width == 0 for d in sheet.column_dimensions.values()
                                   if d.min <= cell.column <= d.max)
                            or not any(left <= cell.column <= right and top <= cell.row <= bottom
                                       for left, top, right, bottom in areas)):
                        raise ValueError('Printed Jordan fields must be visible inside the print area')
                    if (isinstance(cell.value, str) and cell.data_type != 's') or cell.data_type == 'f':
                        raise ValueError('Printed Jordan fields must retain literal readable values')
                    allowed_formats = {'General', '@', '#,##0', '#,##0.####', '#,##0.00', '0.00############'}
                    if cell.coordinate == 'C9':
                        allowed_formats = {'dd-mmm-yyyy'}
                    elif cell.row == total_row and cell.column == 9:
                        allowed_formats = {'"' + data.get('currency', 'USD') + ' "#,##0.00'}
                    if cell.number_format not in allowed_formats:
                        raise ValueError('Printed Jordan fields must retain literal readable formats')
                if sheet['C9'].value is not None and sheet['C9'].number_format != 'dd-mmm-yyyy':
                    raise ValueError('Printed Jordan date must retain its readable date format')
                declarations = data['remarks'].splitlines() if data.get('remarks') is not None else ('NO COMMERCIAL VALUE', 'INVOICE ONLY FOR CUSTOMS PURPOSE')
                for declaration in declarations:
                    cells = [cell for row in sheet for cell in row if declaration in str(cell.value or '').splitlines()]
                    if len(cells) != 1 or cells[0].row <= total_row or not any(
                            left <= cells[0].column <= right and top <= cells[0].row <= bottom
                            for left, top, right, bottom in areas):
                        raise ValueError('Jordan declarations must appear once below the total inside the print area')
    finally:
        book.close()
