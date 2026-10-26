"""
Excel template populator for customs invoices.

Supports 3 regional blank template formats:
  - BD (Bangladesh / EGMCL / PGCL) -> templates/BD_template.xlsx
  - INDIA (India / TGPL)           -> templates/INDIA_template.xlsx
  - JORDAN (Jordan / EJU1)         -> templates/JORDAN_template.xlsx

Includes safe merged-cell handling, label-search fallback for custom templates,
and Extraction_Audit / Extracted_Items sheets.
"""

from pathlib import Path
from copy import copy
from datetime import date
import math
import re
from tempfile import NamedTemporaryFile
from textwrap import wrap
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Border, PatternFill, Side
from .models import FOCData
from .regions import REGIONS, select_region, template_region


def _safe_write(ws, r: int, c: int, val):
    """Safely write to a cell, redirecting to anchor cell if MergedCell."""
    cell = ws.cell(r, c)
    if type(cell).__name__ == "MergedCell":
        for rng in ws.merged_cells.ranges:
            if cell.coordinate in rng:
                cell = ws.cell(rng.min_row, rng.min_col)
                break
    cell.value = val
    if isinstance(val, str):
        cell.data_type = 's'


def _find_label(ws, label: str):
    """Find a cell containing a case-insensitive label substring."""
    label = label.upper()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and label in str(cell.value).upper():
                return cell
    return None


def _next_writable(ws, cell):
    """Search right, then below, avoiding merged non-anchor cells."""
    merged_nonanchors = {
        c.coordinate for rng in ws.merged_cells.ranges
        for row in ws[rng.coord] for c in row
        if c.coordinate != rng.start_cell.coordinate
    }
    for dc, dr in [(1, 0), (2, 0), (0, 1), (1, 1), (0, 2)]:
        c = ws.cell(cell.row + dr, cell.column + dc)
        if c.coordinate not in merged_nonanchors and (c.value is None or str(c.value).strip() == ""):
            return c
    return ws.cell(cell.row, cell.column + 1)


def _split_address_lines(address: str | None) -> list[str]:
    """Split multi-line or comma-separated address into clean lines."""
    if not address:
        return []
    lines = [l.strip() for l in address.split("\n") if l.strip()]
    if len(lines) == 1 and len(lines[0]) > 60:
        parts = [p.strip() for p in lines[0].split(",") if p.strip()]
        lines = parts
    return lines


def _style_goods_table(ws, region, last_row):
    """Outline logical invoice columns, preserving merged descriptions and values."""
    groups = ((1, 1), (2, 6), (7, 9)) if region == 'BD' else ((1, 1), (2, 4), (5, 5), (6, 7), (8, 9))
    header, start, minimum = 30, 33, 38
    if region == 'JORDAN':
        groups = ((1, 1), (2, 5), (6, 6), (7, 7), (8, 8), (9, 10))
        header, start, minimum = 24, 27, 32
    line = Side(style='thin', color='FFB8C2CC')
    for row in range(header, max(minimum, last_row) + 1):
        for first, last in groups:
            for col in range(first, last + 1):
                cell = ws.cell(row, col)
                if row > last_row:
                    cell.border = Border()
                    continue
                cell.border = Border(left=line if col == first else Side(),
                                     right=line if col == last else Side(),
                                     top=line if row == header else Side(),
                                     bottom=line if row == header or row >= start else Side())
                if row == header:
                    cell.fill = PatternFill('solid', fgColor='FFE9EDF2')
                    font = copy(cell.font)
                    font.bold = True
                    cell.font = font
                alignment = copy(cell.alignment)
                alignment.vertical = 'center' if row == header else 'top'
                alignment.wrap_text = True
                alignment.horizontal = ('center' if row == header else
                                        'right' if col == REGIONS[region]['quantity'] or (region == 'INDIA' and col == 9)
                                        or (region == 'JORDAN' and col in (6, 9)) else 'left')
                cell.alignment = alignment
    ws.row_dimensions[header].height = max(30, ws.row_dimensions[header].height or 0)
    for row in (header + 1, header + 2):
        ws.row_dimensions[row].height = 3


def _populate_bd(ws, data: FOCData):
    """Populate BD (Bangladesh) customs invoice template."""
    def red(cell):
        font = copy(cell.font)
        font.color = 'FFFF0000'
        cell.font = font

    for coordinate in ('C9', 'H11', 'H12', 'I13', 'B11', 'E21', 'A23', 'C25', 'B27', 'G27', 'F28', 'G28', 'H28', 'A29', 'G40'):
        cell = ws[coordinate]
        _safe_write(ws, cell.row, cell.column, None)
    for row in range(12, 18):
        _safe_write(ws, row, 2, None)
    ws.merge_cells('C9:F9')
    ws.merge_cells('H9:I9')
    ws['H9'].alignment=copy(ws['H9'].alignment)
    header_alignment=copy(ws['H9'].alignment)
    header_alignment.wrap_text=True
    ws['H9'].alignment=header_alignment
    ws.row_dimensions[9].height=max(30,ws.row_dimensions[9].height or 0)
    if data.export_date:
        _safe_write(ws, 9, 3, date.fromisoformat(data.export_date[:10]))
        ws['C9'].number_format = 'dd-mmm-yyyy'
    if data.shipping_marks:
        _safe_write(ws, 11, 8, data.shipping_marks)  # H11
    _safe_write(ws, 12, 8, data.destination_city)  # H12
    if data.carton_range:
        ws.merge_cells('H13:I13')
        _safe_write(ws, 13, 8, f'C/NO. {data.carton_range}')
        font = ws['H13'].font
        ws['H13'].value = CellRichText(
            TextBlock(InlineFont(rFont=font.name,sz=font.sz,color='FF000000'),'C/NO. '),
            TextBlock(InlineFont(rFont=font.name,sz=font.sz,color='FFFF0000'),data.carton_range))
        carton_alignment=copy(ws['H13'].alignment)
        carton_alignment.wrap_text=True
        carton_alignment.vertical='top'
        ws['H13'].alignment=carton_alignment
        ws.row_dimensions[13].height=max(30,15*math.ceil((6+len(data.carton_range))/13))

    if data.consignee:
        _safe_write(ws, 11, 2, data.consignee)  # B11
        alignment = copy(ws['B11'].alignment)
        alignment.wrap_text = True
        alignment.vertical = 'top'
        ws['B11'].alignment = alignment
        ws.row_dimensions[11].height = max(ws.row_dimensions[11].height or 19,
                                           15 * math.ceil(len(data.consignee) / 55))
    addr_lines = _split_address_lines(data.consignee_address)
    addr_lines = [re.sub(r'^[\[(]?\s*BIN\s*[:#]?\s*([0-9 -]+?)\s*[)\]]?$',
                         lambda match:'BIN : '+match[1].strip(),line,flags=re.I) for line in addr_lines]
    if len(addr_lines) > 6:
        addr_lines = addr_lines[:5] + ['\n'.join(addr_lines[5:])]
    for row in range(12 + len(addr_lines), 18):
        if row != 13 or not data.carton_range:
            ws.row_dimensions[row].height = 4
    for i, line in enumerate(addr_lines):
        ws.merge_cells(start_row=12+i,start_column=2,end_row=12+i,end_column=6)
        _safe_write(ws, 12 + i, 2, line)
        if re.search(r'\bBIN\b',line,re.I):
            red(ws.cell(12+i,2))
        alignment = copy(ws.cell(12 + i, 2).alignment)
        alignment.wrap_text = True
        alignment.vertical='top'
        ws.cell(12 + i, 2).alignment = alignment
        ws.row_dimensions[12 + i].height = max(ws.row_dimensions[12 + i].height or 15,
                                              15 * sum(max(1, math.ceil(len(s) / 55)) for s in line.splitlines()))

    if data.destination_country:
        _safe_write(ws, 21, 5, data.destination_country.upper())  # E21
    _safe_write(ws, 18, 1, 'SHIPMENT TO: ' + (data.destination_city or data.destination_country or ''))
    if data.invoice_number:
        _safe_write(ws, 25, 3, data.invoice_number)  # C25
    if data.import_cover_note:
        ws.merge_cells('A23:I23')
        _safe_write(ws,23,1,data.import_cover_note)
        ws['A23']._style=copy(ws['A42']._style)
        red(ws['A23'])
        alignment=copy(ws['A23'].alignment)
        alignment.wrap_text=True
        ws['A23'].alignment=alignment
        ws.row_dimensions[23].height=max(18,16*math.ceil(len(data.import_cover_note)/105))
    if data.package_count is not None:
        _safe_write(ws, 27, 2, int(data.package_count))  # B27
    if data.gross_weight_kg is not None:
        _safe_write(ws, 27, 7, float(data.gross_weight_kg))  # G27
        ws['G27'].number_format='0.00############'
    ws.row_dimensions[28].height = 4


    # Preserve the BD design, growing its detail area before totals/footer.
    start_row = 33
    carton_labels=[part.strip() for part in (data.carton_range or '').split(',') if part.strip()]
    if not carton_labels and data.package_count:
        # Rose's invoice uses a bare count; other BD formats use the PKGS label.
        carton_labels=[str(int(data.package_count)) if data.shipping_marks=='ROSE-FOC' else f'{int(data.package_count)} PKGS']
    extra = max(0, max(len(data.items),len(carton_labels)) - 6)
    shifted_merges = [copy(rng) for rng in ws.merged_cells.ranges if rng.min_row >= 39]
    if extra:
        for rng in shifted_merges:
            ws.unmerge_cells(str(rng))
        dimensions = {r:copy(dim) for r,dim in ws.row_dimensions.items() if r>=39}
        ws.insert_rows(39, extra)
        for r,dim in sorted(dimensions.items(),reverse=True):
            ws.row_dimensions[r+extra] = dim
            ws.row_dimensions[r+extra].index = r+extra
        for rng in shifted_merges:
            rng.shift(row_shift=extra)
            ws.merge_cells(str(rng))
    # Remove the duplicate footer already present in the supplied blank template.
    ws.cell(45+extra,1).value = None
    ws.cell(46+extra,1).value = None
    for i, item in enumerate(data.items):
        r = start_row + i
        for col in range(1,10):
            ws.cell(r,col)._style = copy(ws.cell(start_row,col)._style)
        ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6)
        ws.cell(r,2).alignment = copy(ws.cell(start_row,2).alignment)
        alignment = copy(ws.cell(r,2).alignment)
        alignment.wrap_text = True
        alignment.vertical = 'top'
        ws.cell(r,2).alignment = alignment
        line_count = sum(max(1, len(wrap(part, width=55))) for part in item.description.splitlines())
        ws.row_dimensions[r].height = max(24, 15 * line_count + 4)
        _safe_write(ws, r, 2, item.description)
        if item.quantity is not None:
            _safe_write(ws, r, 7, item.quantity)
            ws.cell(r,7).number_format='#,##0' if float(item.quantity).is_integer() else '#,##0.####'
        if item.unit:
            _safe_write(ws, r, 8, item.unit)
        for col in (7,8):
            numeric_alignment=copy(ws.cell(r,col).alignment)
            numeric_alignment.vertical='top'
            ws.cell(r,col).alignment=numeric_alignment

    for index,label in enumerate(carton_labels):
        cell=ws.cell(start_row+index,1)
        cell._style=copy(ws.cell(start_row,1)._style)
        _safe_write(ws,cell.row,1,label)
        red(cell)
        alignment=copy(cell.alignment)
        alignment.wrap_text=True
        alignment.vertical='top'
        cell.alignment=alignment
        ws.row_dimensions[cell.row].height=max(ws.row_dimensions[cell.row].height or 15,
                                               15*math.ceil(len(label)/max(1,ws.column_dimensions['A'].width)))

    tot_row = 40 + extra
    _safe_write(ws,tot_row+1,1,None)
    if data.tin_number:
        ws.merge_cells(start_row=tot_row+1,start_column=1,end_row=tot_row+1,end_column=9)
        _safe_write(ws,tot_row+1,1,'TIN # '+data.tin_number)
        ws.cell(tot_row+1,1)._style=copy(ws.cell(tot_row+2,1)._style)
        red(ws.cell(tot_row+1,1))
        ws.row_dimensions[tot_row+1].height=max(18,ws.row_dimensions[tot_row+1].height or 0)
    _safe_write(ws, tot_row, 6, "TOTAL VALUE:")
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row==tot_row and rng.min_col==7:ws.unmerge_cells(str(rng))
    ws.merge_cells(start_row=tot_row,start_column=7,end_row=tot_row,end_column=9)
    _safe_write(ws, tot_row, 7, data.total_value)
    ws.cell(tot_row,7).number_format = '"USD "0.00'
    total_alignment=copy(ws.cell(tot_row,7).alignment)
    total_alignment.wrap_text=True
    ws.cell(tot_row,7).alignment=total_alignment
    ws.row_dimensions[tot_row].height=max(30,ws.row_dimensions[tot_row].height or 0)
    for r in (tot_row+2,tot_row+3):
        ws.row_dimensions[r].height=max(15,ws.row_dimensions[r].height or 0)
    ws.print_area = f'A1:I{43+extra}'
    ws.print_title_rows = '30:30'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = .4
    ws.page_margins.top = .5
    ws.page_margins.bottom = .55
    ws.page_margins.header = ws.page_margins.footer = .2
    # Excel interprets ampersands in header/footer text as field codes.
    identity = (data.invoice_number or Path(data.source_file).stem).replace('&', '&&')
    ws.HeaderFooter.differentFirst = True
    ws.HeaderFooter.differentOddEven = False
    company = str(ws['A1'].value or ws['E1'].value or '').replace('&', '&&')
    ws.oddHeader.left.text = company
    ws.oddHeader.right.text = identity
    for part in (ws.oddHeader.left, ws.oddHeader.right):
        part.font = 'Arial'
        part.size = 8
        part.color = '606060'
    ws.firstHeader.left.text = ws.firstHeader.center.text = ws.firstHeader.right.text = None
    ws.oddFooter.left.text = identity
    ws.oddFooter.center.text = None
    ws.oddFooter.right.text = 'Page &P of &N'
    for part in (ws.oddFooter.left, ws.oddFooter.right):
        part.font = 'Arial'
        part.size = 8
        part.color = '606060'
    ws.firstFooter = copy(ws.oddFooter)
    _style_goods_table(ws, 'BD', max(33, 32 + max(len(data.items), len(carton_labels))))


def _populate_india(ws, data: FOCData):
    """Preserve India's layout while growing goods before its total and declarations."""
    def fit(cell, text, width=55, minimum=20):
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = 'top'
        cell.alignment = alignment
        lines = sum(max(1, len(wrap(part, width=width))) for part in str(text or '').split('\n'))
        height = max(minimum, 14 * lines + 4)
        if height > 409:
            raise ValueError('India invoice text exceeds the printable height of one Excel row')
        ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 0, height)

    for coordinate in ('C9', 'B11', 'H11', 'H12', 'H13', 'D20', 'A22', 'A23', 'C25', 'B27', 'G27', 'I42'):
        cell = ws[coordinate]
        _safe_write(ws, cell.row, cell.column, None)
    for row in range(12, 17):
        _safe_write(ws, row, 2, None)
    for row in range(33, 39):
        for col in range(1, 10):
            _safe_write(ws, row, col, None)
    if data.export_date:
        _safe_write(ws, 9, 3, date.fromisoformat(data.export_date[:10]))
        ws['C9'].number_format = 'dd-mmm-yyyy'
    for coordinate, value in (('B11', data.consignee), ('H11', data.shipping_marks),
                              ('H12', data.destination_city), ('C25', data.invoice_number)):
        _safe_write(ws, ws[coordinate].row, ws[coordinate].column, value)
        fit(ws[coordinate], value, 22 if coordinate.startswith('H') else 55)
    if data.carton_range:
        _safe_write(ws, 13, 8, 'C/NO. '+data.carton_range)
        fit(ws['H13'], ws['H13'].value, 22)
    address = _split_address_lines(data.consignee_address)
    if len(address) > 5:
        address = address[:4] + ['\n'.join(address[4:])]
    for index, line in enumerate(address):
        cell = ws.cell(12 + index, 2)
        _safe_write(ws, cell.row, 2, line)
        fit(cell, line)
        if re.search(r'\b(?:IEC|PAN|GST|BIN|TIN)\b', line, re.I):
            font = copy(cell.font)
            font.color = 'FFFF0000'
            cell.font = font
    for row in range(12 + len(address), 17):
        if row != 13 or not data.carton_range:
            ws.row_dimensions[row].height = 5
    _safe_write(ws, 20, 4, data.destination_country.upper() if data.destination_country else None)
    fit(ws['D20'], ws['D20'].value)
    # Keep shipment labels visible without treating packing dates or addresses as departure facts.
    for row, label, value, suffix in (
            (17, 'SHIPMENT FROM:          TO: ', data.destination_city, '          BY:'),
            (18, 'ON/ABT.:', None, ''),
            (22, 'PORT OF DISCHARGE: ', data.port_of_discharge, ''),
            (23, 'SHIPMENT TERMS: ', data.shipment_terms, '')):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row, 1)
        font = cell.font
        _safe_write(ws, row, 1, label + (value or '') + suffix)
        if value:
            cell.value = CellRichText(
                TextBlock(InlineFont(rFont=font.name, sz=font.sz, color='FF000000'), label),
                TextBlock(InlineFont(rFont=font.name, sz=font.sz, color='FFFF0000'), value))
            if suffix:
                cell.value.append(TextBlock(InlineFont(rFont=font.name, sz=font.sz, color='FF000000'), suffix))
        fit(cell, str(cell.value), 110)
    _safe_write(ws, 27, 2, data.package_count)
    _safe_write(ws, 27, 7, data.gross_weight_kg)

    layout = REGIONS['INDIA']
    start_row = layout['start']
    carton_labels = [part.strip() for part in (data.carton_range or '').split(',') if part.strip()]
    extra = max(0, max(len(data.items), len(carton_labels)) - 6)
    shifted = [copy(rng) for rng in ws.merged_cells.ranges if rng.min_row >= 39]
    if extra:
        for rng in shifted:
            ws.unmerge_cells(str(rng))
        dimensions = {r: copy(dim) for r, dim in ws.row_dimensions.items() if r >= 39}
        ws.insert_rows(39, extra)
        for r, dimension in sorted(dimensions.items(), reverse=True):
            ws.row_dimensions[r + extra] = dimension
            ws.row_dimensions[r + extra].index = r + extra
        for rng in shifted:
            rng.shift(row_shift=extra)
            ws.merge_cells(str(rng))
    for index in range(max(len(data.items), len(carton_labels))):
        row = start_row + index
        for col in range(1, 10):
            ws.cell(row, col)._style = copy(ws.cell(start_row, col)._style)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 26
        if index < len(data.items):
            item = data.items[index]
            for col, value in ((2, item.description), (5, str(item.hs_code) if item.hs_code is not None else None),
                               (layout['quantity'], item.quantity), (layout['unit'], item.unit), (9, item.amount)):
                _safe_write(ws, row, col, value)
            if item.quantity is not None:
                ws.cell(row, layout['quantity']).number_format = '#,##0' if float(item.quantity).is_integer() else '#,##0.####'
            fit(ws.cell(row, 2), item.description, 46, 26)
            fit(ws.cell(row, 5), item.hs_code, 11, 26)
            fit(ws.cell(row, 7), item.unit, 8, 26)
        if index < len(carton_labels):
            _safe_write(ws, row, 1, carton_labels[index])
            fit(ws.cell(row, 1), carton_labels[index], 13, 26)
            font = copy(ws.cell(row, 1).font)
            font.color = 'FFFF0000'
            ws.cell(row, 1).font = font
    for row in range(start_row + max(len(data.items), len(carton_labels)), 39 + extra):
        ws.row_dimensions[row].height = 5
    total_row = 42 + extra
    _safe_write(ws, total_row, layout['total_column'], 'TOTAL VALUE:')
    _safe_write(ws, total_row, 9, data.total_value)
    ws.cell(total_row, 9).number_format = '"USD "#,##0.00'
    ws.print_area = f'A1:I{49 + extra}'
    ws.print_title_rows = '30:30'
    identity = (data.invoice_number or Path(data.source_file).stem).replace('&', '&&')
    ws.HeaderFooter.differentFirst = True
    ws.HeaderFooter.differentOddEven = False
    ws.firstHeader.left.text = ws.firstHeader.center.text = ws.firstHeader.right.text = None
    ws.oddHeader.left.text = str(ws['A1'].value or '').replace('&', '&&')
    ws.oddHeader.right.text = identity
    ws.oddFooter.left.text = identity
    ws.oddFooter.center.text = None
    ws.oddFooter.right.text = 'Page &P of &N'
    for part in (ws.oddHeader.left, ws.oddHeader.right, ws.oddFooter.left, ws.oddFooter.right):
        part.font = 'Arial'
        part.size = 8
        part.color = '606060'
    ws.firstFooter = copy(ws.oddFooter)
    _style_goods_table(ws, 'INDIA', max(33, 32 + max(len(data.items), len(carton_labels))))


def _populate_jordan(ws, data: FOCData):
    """Grow Jordan goods while preserving separate KG, quantity, unit and value."""
    def fit(cell, value, width, minimum=20):
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = 'top'
        cell.alignment = alignment
        lines = sum(max(1, len(wrap(part, width=width))) for part in str(value or '').split('\n'))
        height = max(minimum, 14 * lines + 4)
        if height > 409:
            raise ValueError('Jordan invoice text exceeds the printable height of one Excel row')
        ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 0, height)

    for coordinate, value, width in (
            ('C9', date.fromisoformat(data.export_date[:10]) if data.export_date else None, 28),
            ('I9', data.shipping_marks, 14), ('I10', data.destination_city, 14),
            ('I11', data.carton_range, 14), ('C11', data.consignee, 58),
            ('D16', data.destination_country.upper() if data.destination_country else None, 65),
            ('C18', data.invoice_number, 75), ('B21', data.package_count, 12),
            ('F21', data.gross_weight_kg, 10)):
        cell = ws[coordinate]
        _safe_write(ws, cell.row, cell.column, value)
        fit(cell, value, width)
    ws['C9'].number_format = 'dd-mmm-yyyy'
    alignment = copy(ws['C9'].alignment)
    alignment.horizontal = 'left'
    ws['C9'].alignment = alignment
    ws['F21'].number_format = '0.00############'
    address = _split_address_lines(data.consignee_address)
    if len(address) > 4:
        address = address[:3] + ['\n'.join(address[3:])]
    for index in range(4):
        value = address[index] if index < len(address) else None
        _safe_write(ws, 12 + index, 3, value)
        fit(ws.cell(12 + index, 3), value, 58, 5 if value is None else 20)

    start = REGIONS['JORDAN']['start']
    cartons = [part.strip() for part in (data.carton_range or '').split(',') if part.strip()]
    count = max(len(data.items), len(cartons))
    extra = max(0, count - 6)
    if extra:
        shifted = [copy(rng) for rng in ws.merged_cells.ranges if rng.min_row >= 33]
        for rng in shifted:
            ws.unmerge_cells(str(rng))
        dimensions = {r: copy(dim) for r, dim in ws.row_dimensions.items() if r >= 33}
        ws.insert_rows(33, extra)
        for row, dimension in sorted(dimensions.items(), reverse=True):
            ws.row_dimensions[row + extra] = dimension
            ws.row_dimensions[row + extra].index = row + extra
        for rng in shifted:
            rng.shift(row_shift=extra)
            ws.merge_cells(str(rng))
    for index in range(max(6, count)):
        row = start + index
        for col in range(1, 13):
            ws.cell(row, col)._style = copy(ws.cell(start, col)._style)
            _safe_write(ws, row, col, None)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        ws.row_dimensions[row].height = 26 if index < count else 5
        if index < len(cartons):
            _safe_write(ws, row, 1, cartons[index])
            fit(ws.cell(row, 1), cartons[index], 11, 26)
            font = copy(ws.cell(row, 1).font)
            font.color = 'FFFF0000'
            ws.cell(row, 1).font = font
        if index < len(data.items):
            item = data.items[index]
            for col, value in ((2, item.description), (6, item.weight_kg), (7, item.quantity),
                               (8, item.unit), (9, item.amount), (11, item.unit_price)):
                _safe_write(ws, row, col, value)
            if item.unit_price is not None:
                _safe_write(ws, row, 12, data.currency)
            fit(ws.cell(row, 2), item.description, 52, 26)
            fit(ws.cell(row, 8), item.unit, 13, 26)
            ws.cell(row, 6).number_format = '0.00############'
            ws.cell(row, 7).number_format = '#,##0' if item.quantity is not None and float(item.quantity).is_integer() else '#,##0.####'
            ws.cell(row, 9).number_format = '#,##0.00'
    total = 35 + extra
    _safe_write(ws, 24, 9, f'VALUE ({data.currency})')
    _safe_write(ws, total, 8, 'TOTAL VALUE:')
    _safe_write(ws, total, 9, data.total_value)
    ws.cell(total, 9).number_format = '"' + data.currency + ' "#,##0.00'
    ws.print_area = f'A1:J{39 + extra}'
    ws.print_title_rows = '24:24'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.45
    ws.page_margins.header = ws.page_margins.footer = 0.2
    identity = (data.invoice_number or Path(data.source_file).stem).replace('&', '&&')
    ws.HeaderFooter.differentFirst = True
    ws.HeaderFooter.differentOddEven = False
    ws.firstHeader.left.text = ws.firstHeader.center.text = ws.firstHeader.right.text = None
    ws.oddHeader.left.text = str(ws['A1'].value or '').replace('&', '&&')
    ws.oddHeader.right.text = identity
    ws.oddFooter.left.text = identity
    ws.oddFooter.center.text = None
    ws.oddFooter.right.text = 'Page &P of &N'
    for part in (ws.oddHeader.left, ws.oddHeader.right, ws.oddFooter.left, ws.oddFooter.right):
        part.font = 'Arial'
        part.size = 8
        part.color = '606060'
    ws.firstFooter = copy(ws.oddFooter)
    _style_goods_table(ws, 'JORDAN', max(start, start + count - 1))


def populate_template(template: Path, output: Path, data: FOCData):
    """
    Populate a customs invoice template with extracted FOC data.
    """
    selected = select_region(data)
    if template_region(template) not in (None, selected):
        raise ValueError('Template does not match the selected invoice region')
    wb = load_workbook(template)
    ws = wb.active

    if selected == "BD":
        _populate_bd(ws, data)
    elif selected == "INDIA":
        _populate_india(ws, data)
    elif selected == "JORDAN":
        _populate_jordan(ws, data)
    else:
        # Fallback to label search
        labels = {
            "DATE OF EXPORTATION": "export_date",
            "CUSTOMS INVOICE NO": "invoice_number",
            "COUNTRY OF ULTIMATE DESTINATION": "destination_country",
            "TOTAL VALUE": "total_value",
            "G.W": "gross_weight_kg",
            "GW": "gross_weight_kg",
            "TOTAL": "package_count",
        }
        for label, attr in labels.items():
            src = _find_label(ws, label)
            val = getattr(data, attr, None)
            if src and val is not None:
                target = _next_writable(ws, src)
                _safe_write(ws, target.row, target.column, val)

        desc = _find_label(ws, "DESCRIPTION  OF  GOODS") or _find_label(ws, "DESCRIPTION OF GOODS")
        qty = _find_label(ws, "QUANTITY")
        hs = _find_label(ws, "HS CODE")
        if desc and qty:
            start = max(desc.row, qty.row) + 1
            for i, item in enumerate(data.items[:10]):
                r = start + i
                _safe_write(ws, r, desc.column, item.description)
                _safe_write(ws, r, qty.column, item.quantity)
                if qty.column + 1 <= ws.max_column:
                    _safe_write(ws, r, qty.column + 1, item.unit)
                if hs and item.hs_code:
                    _safe_write(ws, r, hs.column, item.hs_code)

    # Add audit sheet for traceability and compliance
    audit = wb.create_sheet("Extraction_Audit")
    audit.append(["Field", "Extracted Value"])
    for k, v in data.to_dict().items():
        if k != "items" and not (selected == 'BD' and k == 'net_weight_kg'):
            audit.append([k, str(v) if v is not None else ""])

    if data.items:
        items_sheet = wb.create_sheet("Extracted_Items")
        items_sheet.append(["Index", "Description", "Quantity", "Unit", "Weight_KG", "HS_Code", "Unit_Price", "Amount", "Source_Page", "Confidence", "Source_Rows"])
        for idx, it in enumerate(data.items, 1):
            items_sheet.append([
                idx, it.description, it.quantity, it.unit,
                it.weight_kg, it.hs_code, it.unit_price, it.amount, it.source_page, it.confidence, '; '.join(it.source_refs),
            ])
            items_sheet.cell(idx + 1, 7).number_format = '0.00####'
            items_sheet.cell(idx + 1, 8).number_format = '0.00'
        items_sheet.column_dimensions['B'].width = 90
        items_sheet.freeze_panes = 'A2'
    audit.column_dimensions['A'].width = 25
    audit.column_dimensions['B'].width = 100
    audit.freeze_panes = 'A2'
    for sheet in wb.worksheets:
        if sheet is ws:
            continue
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'

    Path(output).parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=Path(output).parent, suffix='.xlsx', delete=False) as handle:
        temporary = Path(handle.name)
    try:
        wb.save(temporary)
        temporary.replace(output)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
