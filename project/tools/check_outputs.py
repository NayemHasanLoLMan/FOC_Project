"""Check generated sample invoices against their source page coverage and audit."""
import sys,json
from pathlib import Path
import openpyxl,pymupdf
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from tools.compare_invoices import read_invoice
records=[]
files=sorted((ROOT/'output').glob('*.xlsx'))
if not files:raise SystemExit('No sample invoices found in project/output')
for path in files:
    wb=openpyxl.load_workbook(path)
    audit={r[0]:r[1] for r in wb['Extraction_Audit'].iter_rows(values_only=True)}
    occupied=set()
    for merged in wb.active.merged_cells.ranges:
        assert not occupied.intersection(merged.cells),path.name
        occupied.update(merged.cells)
    with pymupdf.open(ROOT/'examples'/audit['source_file']) as pdf: actual_pages=len(pdf)
    parsed=read_invoice(path)
    assert parsed['audit_matches_invoice'],path.name
    assert int(audit['pages_processed'])==actual_pages,path.name
    assert int(audit['pages_total'])==actual_pages,path.name
    assert not any(str(c.value).startswith('#REF!') for row in wb.active for c in row),path.name
    assert not any(any(term in str(c.value).upper() for term in ('REVIEW REQUIRED','DRAFT','NEEDS VALIDATION','HUMAN INTERVENTION')) for row in wb.active for c in row),path.name
    assert wb.active['A29'].value is None,path.name
    for field,coordinate in [('export_date','C9'),('gross_weight_kg','G27'),('package_count','B27')]:
        if audit.get(field) is None:assert wb.active[coordinate].value is None,(path.name,field)
    total=next(c for row in wb.active for c in row if c.value=='TOTAL VALUE:')
    if audit.get('total_value') is None:assert wb.active.cell(total.row,7).value is None,path.name
    assert total.row>32+len(parsed['items']),path.name
    records.append({'output':path.name,'pages':actual_pages,'items':len(parsed['items']),'audit_matches_invoice':True})
    wb.close()
result={'files':len(records),'pages':sum(r['pages'] for r in records),'items':sum(r['items'] for r in records),'records':records}
(ROOT/'reports').mkdir(exist_ok=True)
(ROOT/'reports/validation.json').write_text(json.dumps(result,indent=2))
print(json.dumps({k:v for k,v in result.items() if k!='records'},indent=2))
